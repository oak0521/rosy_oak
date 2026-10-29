import pandas as pd, numpy as np, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L
from openpyxl.formatting.rule import FormulaRule

import os
HERE = os.path.dirname(os.path.abspath(__file__))
U = os.path.join(HERE, '..', 'source') + '/'
OUT = os.path.join(HERE, '..', 'build') + '/'
os.makedirs(OUT, exist_ok=True)
ad = pd.read_pickle(OUT + 'ad_agg.pkl')
ssd = pd.read_pickle(OUT + 'ss_daily.pkl')
ssp = pd.read_pickle(OUT + 'ss_prod.pkl')

FONT = 'Arial'
NAVY, LIGHT, INPUTFILL, GREY = '1F3864', 'D9E1F2', 'FFF2CC', 'F2F2F2'
f_norm = Font(name=FONT, size=10)
f_bold = Font(name=FONT, size=10, bold=True)
f_hdr = Font(name=FONT, size=10, bold=True, color='FFFFFF')
f_title = Font(name=FONT, size=16, bold=True, color=NAVY)
f_sec = Font(name=FONT, size=12, bold=True, color=NAVY)
f_in = Font(name=FONT, size=10, color='0000FF')
f_note = Font(name=FONT, size=9, color='595959')
fill_hdr = PatternFill('solid', fgColor=NAVY)
fill_sub = PatternFill('solid', fgColor=LIGHT)
fill_in = PatternFill('solid', fgColor=INPUTFILL)
fill_grey = PatternFill('solid', fgColor=GREY)
thin = Side(style='thin', color='BFBFBF')
box = Border(left=thin, right=thin, top=thin, bottom=thin)
CEN = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT = Alignment(horizontal='left', vertical='center', wrap_text=True)
LEFTTOP = Alignment(horizontal='left', vertical='top', wrap_text=True)
RIGHT = Alignment(horizontal='right', vertical='center')

WON, MAN, PCT, PCT1, INT = '#,##0', '#,##0.0', '0%', '0.0%', '#,##0'
PCTD = '0%;-0%;"-"'

wb = Workbook()
wb.remove(wb.active)
def sheet(name, tab=None):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    if tab: ws.sheet_properties.tabColor = tab
    return ws

def put(ws, ref, v, font=f_norm, fmt=None, fill=None, al=None, border=False):
    c = ws[ref]
    c.value = v
    c.font = font
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    if al: c.alignment = al
    if border: c.border = box
    return c

def hdr(ws, row, c0, labels, fill=fill_hdr, font=f_hdr, height=None):
    for i, t in enumerate(labels):
        put(ws, f'{L(c0+i)}{row}', t, font, fill=fill, al=CEN, border=True)
    if height: ws.row_dimensions[row].height = height

def widths(ws, d):
    for k, v in d.items():
        ws.column_dimensions[k].width = v

S_GA = "'가정'"; S_AD = "'데이터_광고'"; S_SD = "'데이터_SS일별'"; S_SP = "'데이터_SS상품'"
S_CK = "'점검'"; S_TV = "'시간제한_검증'"; S_JM = "'자사몰_매출'"; S_DJ = "'데이터_자사몰'"; S_WK = "'주차별_추이'"; S_RS = "'10월_실적'"; S_TG = "'목표_11-12월'"; S_AN = "'성과분석'"

BRANDS = ['아넬라', '로지오가닉']
MEDIA7 = ['파워링크', '쇼핑검색', '브랜드검색', 'GFA 카탈로그', 'GFA 전환', '메타 스스', '메타 자사몰']
MEDIA6 = [m for m in MEDIA7 if m != '브랜드검색']

# =====================================================================
# 가정
# =====================================================================
ga = sheet('가정', '7F7F7F')
widths(ga, {'A': 30, 'B': 18, 'C': 14, 'D': 14, 'E': 14, 'F': 70})
put(ga, 'A1', '가정·기준값', f_title)
put(ga, 'A2', '파란 글씨·노란 칸 = 입력값(수정 가능), 검정 = 수식. 다른 시트의 계산은 모두 이 시트와 데이터 시트 참조.', f_note)

put(ga, 'A4', '기간', f_sec)
rows = [(5, '집계 시작일', dt.date(2026, 10, 1), 'yyyy-mm-dd', True, '광고 리포트·스마트스토어 파일의 시작일'),
        (6, '집계 종료일', dt.date(2026, 10, 28), 'yyyy-mm-dd', True, '광고 리포트·스마트스토어 파일의 종료일'),
        (7, '집계 일수', '=B6-B5+1', INT, False, ''),
        (8, '10월 일수', 31, INT, True, ''),
        (9, '10월 잔여일수', '=B8-B7', INT, False, '월말 예상치 계산에 사용 (10/29~10/31)'),
        (10, '11월 일수', 30, INT, True, ''),
        (11, '12월 일수', 31, INT, True, ''),
        (12, '"근접" 판정 기준 (달성률)', 0.95, PCT, True, '달성률 100% 이상 = 달성, 이 값 이상 = 근접, 미만 = 미달')]
for r, a, b, fmt, inp, note in rows:
    put(ga, f'A{r}', a, f_norm, border=True)
    put(ga, f'B{r}', b, f_in if inp else f_norm, fmt, fill_in if inp else None, RIGHT, True)
    put(ga, f'F{r}', note, f_note)

put(ga, 'A14', '월 광고 예산 (만원, 네이버 VAT 미포함)', f_sec)
hdr(ga, 15, 1, ['브랜드', '매체', '10월', '11월', '12월', '비고'])
budget = [
    ('아넬라', '파워링크', 120, 120, ''),
    ('아넬라', '쇼핑검색', 380, 380, ''),
    ('아넬라', 'GFA 카탈로그', 30, 30, ''),
    ('아넬라', 'GFA 전환', 250, 150, '신규 소구 5종 테스트가 소재 검토로 HOLD 중이라 11월 감액'),
    ('아넬라', '메타 스스', 200, 200, ''),
    ('아넬라', '메타 자사몰', 250, 180, '웰컴박스·기획전 보류로 축소'),
    ('로지오가닉', '파워링크', 120, 120, ''),
    ('로지오가닉', '쇼핑검색', 250, 250, ''),
    ('로지오가닉', 'GFA 카탈로그', 30, 30, ''),
    ('로지오가닉', 'GFA 전환', 0, 20, '11월 신설, 넾다세일 한정 예산. 12월은 11월 유지 가정에 따라 같은 값이 들어가므로 사용처 확정 필요'),
    ('로지오가닉', '메타 스스', 250, 215, '넾다세일로 약간 감액'),
    ('로지오가닉', '메타 자사몰', 180, 150, '보리차 시크릿링크 보류로 감액'),
]
for i, (b, m, o, n, note) in enumerate(budget):
    r = 16 + i
    put(ga, f'A{r}', b, border=True); put(ga, f'B{r}', m, border=True)
    put(ga, f'C{r}', o, f_in, INT, fill_in, RIGHT, True)
    put(ga, f'D{r}', n, f_in, INT, fill_in, RIGHT, True)
    put(ga, f'E{r}', f'=D{r}', f_norm, INT, None, RIGHT, True)
    put(ga, f'F{r}', note, f_note, al=LEFT)
put(ga, 'F15', '비고 (12월은 "11월 예산 유지" 가정으로 11월 값을 따라가며, 다른 값이 필요하면 E열을 덮어쓰기)', f_hdr, fill=fill_hdr, al=CEN, border=True)
for r, b in [(28, '아넬라'), (29, '로지오가닉')]:
    put(ga, f'A{r}', b, f_bold, fill=fill_sub, border=True); put(ga, f'B{r}', '합계', f_bold, fill=fill_sub, border=True)
    for c in 'CDE':
        put(ga, f'{c}{r}', f'=SUMIFS({c}$16:{c}$27,$A$16:$A$27,$A{r})', f_bold, INT, fill_sub, RIGHT, True)
put(ga, 'A30', '전체', f_bold, fill=fill_sub, border=True); put(ga, 'B30', '합계', f_bold, fill=fill_sub, border=True)
for c in 'CDE':
    put(ga, f'{c}30', f'={c}28+{c}29', f_bold, INT, fill_sub, RIGHT, True)
put(ga, 'A32', '제시 합계(검증용) 아넬라', border=True)
put(ga, 'A33', '제시 합계(검증용) 로지오가닉', border=True)
for r, (o, n) in [(32, (1230, 1060)), (33, (830, 785))]:
    put(ga, f'C{r}', o, f_in, INT, fill_in, RIGHT, True); put(ga, f'D{r}', n, f_in, INT, fill_in, RIGHT, True)
    put(ga, f'E{r}', f'=D{r}', f_norm, INT, None, RIGHT, True)
put(ga, 'F32', '출처: 사용자가 전달한 10월·11월 예산 총액. 12월은 11월 유지', f_note)
put(ga, 'A34', '검증: 합계 차이 (0이면 일치)', f_bold, border=True)
for c in 'CDE':
    put(ga, f'{c}34', f'=ABS({c}28-{c}32)+ABS({c}29-{c}33)', f_bold, INT, None, RIGHT, True)

put(ga, 'A36', '브랜드검색 정액 (예산표 외 별도 과금)', f_sec)
put(ga, 'A37', '월 정액 (원, 브랜드별)', border=True)
put(ga, 'B37', 880000, f_in, WON, fill_in, RIGHT, True)
put(ga, 'F37', '출처: 9월 월간 리포트 하단 주석 (정액제 2,640,000원/90일을 30일분으로 환산). 주간 리포트의 광고비 집계와 같은 방식', f_note)
put(ga, 'A38', '일 안분액 (원)', border=True)
put(ga, 'B38', '=B37/30', f_norm, WON, None, RIGHT, True)

put(ga, 'A40', '4분기 KPI 보고안 기준 (2026.10.06 갱신본)', f_sec)
hdr(ga, 41, 1, ['항목', '아넬라', '로지오가닉', '', '', '출처·비고'])
kpi = [(42, '광고 ROAS 하한', 8.5, 4.2, PCT, '4분기 KPI 보고안 "ROAS 하한 근거" (8·9월 월간 ROAS가 이 값 이상)'),
       (43, '광고 기준 목표 주문 (건/월)', 2700, 1700, INT, '4분기 KPI 보고안 (9월 ROAS 유지 시 월 환산)'),
       (44, 'KPI안 광고예산 10월 (만원)', 1300, 900, INT, '실행 예산 + 브랜드검색 정액(월 88만원)이 포함된 값'),
       (45, 'KPI안 광고예산 11월 (만원)', 1200, 800, INT, ''),
       (46, 'KPI안 광고예산 12월 (만원)', 1300, 800, INT, '')]
for r, a, x, y, fmt, note in kpi:
    put(ga, f'A{r}', a, border=True)
    put(ga, f'B{r}', x, f_in, fmt, fill_in, RIGHT, True); put(ga, f'C{r}', y, f_in, fmt, fill_in, RIGHT, True)
    put(ga, f'F{r}', note, f_note)

put(ga, 'A48', '기간 구분 (광고 성과를 상시/행사로 나누는 기준)', f_sec)
hdr(ga, 49, 1, ['기간 코드', '시작일', '종료일', '일수', '구분', '주요 이슈'])
periods = [('P1', dt.date(2026, 10, 1), dt.date(2026, 10, 4), '상시', '9주차(9/28~10/4) 중 10/1~10/4만 포함'),
           ('P2', dt.date(2026, 10, 5), dt.date(2026, 10, 11), '상시', '10월 2주차. 4주차 주간 리포트가 "상시 기간"으로 언급'),
           ('P3', dt.date(2026, 10, 12), dt.date(2026, 10, 18), '행사', '아넬라: 라이프위크·오늘끝딜·1+1 이벤트 (3주차 주간 리포트). 로지도 라이프위크'),
           ('P4', dt.date(2026, 10, 19), dt.date(2026, 10, 25), '상시', '1+1은 10/19 오전 종료 후 별도 행사 없음 (4주차 주간 리포트)'),
           ('P5', dt.date(2026, 10, 26), dt.date(2026, 10, 28), '행사', '10/26 넾다세일 시작 (아넬라·로지 모두 11/8까지)')]
for i, (c, s, e, k, note) in enumerate(periods):
    r = 50 + i
    put(ga, f'A{r}', c, border=True)
    put(ga, f'B{r}', s, f_in, 'yyyy-mm-dd', fill_in, RIGHT, True)
    put(ga, f'C{r}', e, f_in, 'yyyy-mm-dd', fill_in, RIGHT, True)
    put(ga, f'D{r}', f'=C{r}-B{r}+1', f_norm, INT, None, RIGHT, True)
    put(ga, f'E{r}', k, f_in, None, fill_in, CEN, True)
    put(ga, f'F{r}', note, f_note, al=LEFT)
ga.freeze_panes = 'A3'
PROW = {p[0]: 50 + i for i, p in enumerate(periods)}   # P1->50 ... P5->54

# =====================================================================
# 데이터_광고
# =====================================================================
da = sheet('데이터_광고', 'A6A6A6')
widths(da, {'A': 12, 'B': 14, 'C': 46, 'D': 8, 'E': 14, 'F': 12, 'G': 10, 'H': 10, 'I': 14, 'J': 8})
hdr(da, 1, 1, ['브랜드', '매체', '캠페인·광고그룹', '기간', '광고비(원)', '노출', '클릭', '구매', '매출(원)', '구분'], height=26)
for i, r in ad.iterrows():
    x = 2 + i
    put(da, f'A{x}', r.brand); put(da, f'B{x}', r.media); put(da, f'C{x}', r['name']); put(da, f'D{x}', r.period)
    if r.media == '브랜드검색':
        put(da, f'E{x}', f"={S_GA}!$B$38*INDEX({S_GA}!$D$50:$D$54,MATCH(D{x},{S_GA}!$A$50:$A$54,0))", f_norm, WON)
    else:
        put(da, f'E{x}', float(r.cost), f_in, WON)
    put(da, f'F{x}', float(r.imp), f_in, WON); put(da, f'G{x}', float(r.clk), f_in, WON)
    put(da, f'H{x}', float(r.buy), f_in, WON); put(da, f'I{x}', float(r.rev), f_in, WON)
    put(da, f'J{x}', f"=INDEX({S_GA}!$E$50:$E$54,MATCH(D{x},{S_GA}!$A$50:$A$54,0))")
AD_LAST = 1 + len(ad)
da.freeze_panes = 'A2'
da['A1'].comment = Comment(
    'Source: 네이버 검색광고 일별 리포트(1948280, 2026.10.01~10.28), 네이버 GFA result.csv, 메타 광고관리자 Raw Data Report 2건(자사몰·스마트스토어). '
    '일 단위 원본을 캠페인·광고그룹×기간(P1~P5)으로 합산해 입력, 사용자 업로드, 2026-10-29. '
    '네이버 검색광고·GFA는 구매완료 전환수·전환매출, 메타는 구매·구매 전환값 기준. '
    '브랜드검색은 리포트상 비용이 0이어서 정액을 일할 안분한 수식 값이며 클릭·노출·구매·매출은 리포트 값.', 'Claude')
da['A1'].comment.width = 420; da['A1'].comment.height = 200

# =====================================================================
# 데이터_SS일별
# =====================================================================
sd = sheet('데이터_SS일별', 'A6A6A6')
widths(sd, {'A': 12, 'B': 12, 'C': 8, 'D': 12, 'E': 16, 'F': 16, 'G': 14})
hdr(sd, 1, 1, ['스토어', '날짜', '구분', '결제건수', '판매금액(총,원)', '판매금액(순,원)', '환불금액(원)'], height=26)
ssd = ssd.sort_values(['store', 'date', 'kind']).reset_index(drop=True)
for i, r in ssd.iterrows():
    x = 2 + i
    put(sd, f'A{x}', r.store); put(sd, f'B{x}', r.date.date(), f_in, 'yyyy-mm-dd'); put(sd, f'C{x}', r.kind)
    put(sd, f'D{x}', float(r['상품결제건수']), f_in, WON); put(sd, f'E{x}', float(r['판매금액(총)']), f_in, WON)
    put(sd, f'F{x}', float(r['판매금액(순)']), f_in, WON); put(sd, f'G{x}', float(r['환불금액']), f_in, WON)
sd.freeze_panes = 'A2'
sd['A1'].comment = Comment(
    'Source: 스마트스토어 통계 SALES 리포트 2건(sales_20261001-20261028, 스토어별), 사용자 업로드, 2026-10-29. '
    '상품 단위 일별 행을 스토어×날짜×구분으로 합산. 구분은 상품명에 "공동구매"가 들어가면 "공구", 아니면 "일반". '
    '공구 상품은 요청에 따라 일반 상품 집계에서 제외하고 참고용으로만 남김.', 'Claude')
sd['A1'].comment.width = 420; sd['A1'].comment.height = 180
SD_LAST = 1 + len(ssd)

# =====================================================================
# 데이터_SS상품
# =====================================================================
sp = sheet('데이터_SS상품', 'A6A6A6')
widths(sp, {'A': 12, 'B': 16, 'C': 70, 'D': 12, 'E': 16, 'F': 16, 'G': 14, 'H': 8})
hdr(sp, 1, 1, ['스토어', '채널상품번호', '채널상품명', '결제건수', '판매금액(총,원)', '판매금액(순,원)', '환불금액(원)', '구분'], height=26)
for i, r in ssp.reset_index(drop=True).iterrows():
    x = 2 + i
    put(sp, f'A{x}', r.store); put(sp, f'B{x}', "'" + str(r['채널상품번호']) if False else str(r['채널상품번호']))
    put(sp, f'C{x}', r['채널상품명'])
    put(sp, f'D{x}', float(r['상품결제건수']), f_in, WON); put(sp, f'E{x}', float(r['판매금액(총)']), f_in, WON)
    put(sp, f'F{x}', float(r['판매금액(순)']), f_in, WON); put(sp, f'G{x}', float(r['환불금액']), f_in, WON)
    put(sp, f'H{x}', f'=IF(ISNUMBER(SEARCH("공동구매",C{x})),"공구","일반")')
    sp[f'B{x}'].number_format = '@'
sp.freeze_panes = 'A2'
SP_LAST = 1 + len(ssp)
sp['A1'].comment = Comment(
    'Source: 스마트스토어 통계 SALES 리포트 2건(sales_20261001-20261028), 사용자 업로드, 2026-10-29. '
    '상품별 28일 합계. 구분(H열)은 상품명에 "공동구매"가 포함되면 "공구"로 표시하는 수식.', 'Claude')
sp['A1'].comment.width = 400; sp['A1'].comment.height = 140

def adr(col):  # data_광고 range
    return f"{S_AD}!${col}$2:${col}$500"
def sdr(col):
    return f"{S_SD}!${col}$2:${col}$400"
def spr(col):
    return f"{S_SP}!${col}$2:${col}$600"

# =====================================================================
# 점검
# =====================================================================
ck = sheet('점검', 'C00000')
widths(ck, {'A': 26, 'B': 18, 'C': 14, 'D': 16, 'E': 22, 'F': 20, 'G': 60})
put(ck, 'A1', '데이터 점검', f_title)
put(ck, 'A2', '팀장 확인 필요 데이터 이슈 및 목표 산정용 보정값.', f_note)

put(ck, 'A4', '1. 로지오가닉 메타 스스: 10/24~10/25 전환이 스토어 실제 결제보다 큼', f_sec)
hdr(ck, 5, 1, ['날짜', '메타 광고비(원)', '메타 구매(건)', '메타 구매 전환값(원)', '로지 스토어 결제금액(총, 공구 포함, 원)', '메타 전환값 ÷ 스토어 결제금액', '보정 대상 (Y = 목표 산정에서 제외)'], height=44)
anom = [(dt.date(2026, 10, 22), 60129, 20, 403681, 'N'), (dt.date(2026, 10, 23), 68710, 17, 428400, 'N'),
        (dt.date(2026, 10, 24), 53199, 70, 2670512, 'Y'), (dt.date(2026, 10, 25), 80781, 108, 3970440, 'Y'),
        (dt.date(2026, 10, 26), 18966, 46, 1839198, 'N')]
for i, (d, c, b, v, flag) in enumerate(anom):
    r = 6 + i
    put(ck, f'A{r}', d, f_in, 'yyyy-mm-dd', al=CEN, border=True)
    put(ck, f'B{r}', c, f_in, WON, border=True); put(ck, f'C{r}', b, f_in, WON, border=True); put(ck, f'D{r}', v, f_in, WON, border=True)
    put(ck, f'E{r}', f'=SUMIFS({sdr("E")},{sdr("A")},"로지오가닉",{sdr("B")},A{r})', f_norm, WON, border=True)
    put(ck, f'F{r}', f'=IF(E{r}=0,0,D{r}/E{r})', f_norm, PCT, border=True)
    put(ck, f'G{r}', flag, f_in, None, fill_in, CEN, True)
ck['A6'].comment = Comment('Source: 메타 광고관리자 Raw Data Report(스마트스토어 캠페인 SUS_로지오가닉_과자(상시)_구매_260506), 사용자 업로드, 2026-10-29. 일별 값 그대로 입력.', 'Claude')
ck['A6'].comment.width = 380; ck['A6'].comment.height = 90
put(ck, 'A12', '보정값 (목표 산정용)', f_bold)
hdr(ck, 13, 1, ['브랜드', '매체', '', '보정 광고비(원)', '보정 매출(원)'])
put(ck, 'A14', '로지오가닉', border=True); put(ck, 'B14', '메타 스스', border=True)
put(ck, 'D14', '=SUMIFS(B6:B10,G6:G10,"Y")', f_norm, WON, border=True)
put(ck, 'E14', '=SUMIFS(D6:D10,G6:G10,"Y")', f_norm, WON, border=True)
put(ck, 'A15', '이틀 합계 전환값이 같은 이틀 스토어 결제금액(총)의 몇 배인가', f_note)
put(ck, 'D15', '=IF(SUMIFS(E6:E10,G6:G10,"Y")=0,0,E14/SUMIFS(E6:E10,G6:G10,"Y"))', f_bold, '0.0"배"')
put(ck, 'A16', '해석: 같은 이틀 로지 스토어 전체 결제(공구 포함)보다 메타 보고 전환값이 큼. 실제 주문이 아닌 전환 집계 이슈(협력광고 연동·중복 집계 등) 가능성. 메타·네이버 협력광고 담당 확인 필요. 11~12월 목표 산정에서는 이 이틀 제외(목표 시트의 보정 스위치).', f_note, al=LEFTTOP)
ck.merge_cells('A16:G17'); ck.row_dimensions[16].height = 30

put(ck, 'A19', '2. 광고 귀속 매출과 스마트스토어 일반상품 순매출 비교', f_sec)
hdr(ck, 20, 1, ['스토어(브랜드)', '스마트스토어 일반상품 순매출(원)', '스스향 광고 귀속 매출(원)', '귀속 매출 ÷ 순매출', '공구 상품 판매(총, 참고, 원)', '공구 비중 (스토어 결제 총액 대비)', '메모'], height=44)
for i, b in enumerate(BRANDS):
    r = 21 + i
    wr = 4 + 21 * i
    put(ck, f'A{r}', b, border=True)
    put(ck, f'B{r}', f"={S_WK}!G{wr+13}", f_norm, WON, border=True)
    put(ck, f'C{r}', f"={S_WK}!G{wr+6}-SUMIFS({adr('I')},{adr('A')},A{r},{adr('B')},\"메타 자사몰\")", f_norm, WON, border=True)
    put(ck, f'D{r}', f'=IF(B{r}=0,0,C{r}/B{r})', f_norm, PCT, border=True)
    put(ck, f'E{r}', f'=SUMIFS({sdr("E")},{sdr("A")},A{r},{sdr("C")},"공구")', f_norm, WON, border=True)
    put(ck, f'F{r}', f'=IF(E{r}+SUMIFS({sdr("E")},{sdr("A")},A{r},{sdr("C")},"일반")=0,0,E{r}/(E{r}+SUMIFS({sdr("E")},{sdr("A")},A{r},{sdr("C")},"일반")))', f_norm, PCT, border=True)
put(ck, 'G21', '아넬라 스토어는 공동구매 상품 없음.', f_note)
put(ck, 'G22', '순매출 대비 귀속 매출이 100%에 가까우면 일반상품 매출 대부분이 광고 귀속. 귀속 기준 차이로 100% 초과 가능.', f_note, al=LEFT)
ck.row_dimensions[22].height = 36

put(ck, 'A25', '3. 데이터 정합 점검 (0이면 일치)', f_sec)
hdr(ck, 26, 1, ['스토어', '일별 시트 일반 순매출', '상품 시트 일반 순매출', '차이'])
for i, b in enumerate(BRANDS):
    r = 27 + i
    put(ck, f'A{r}', b, border=True)
    put(ck, f'B{r}', f'=SUMIFS({sdr("F")},{sdr("A")},A{r},{sdr("C")},"일반")', f_norm, WON, border=True)
    put(ck, f'C{r}', f'=SUMIFS({spr("F")},{spr("A")},A{r},{spr("H")},"일반")', f_norm, WON, border=True)
    put(ck, f'D{r}', f'=B{r}-C{r}', f_bold, WON, border=True)

put(ck, 'A31', '4. GFA 귀속 지연 가능성 (광고비 0원인데 매출이 잡힌 그룹)', f_sec)
hdr(ck, 32, 1, ['그룹', '광고비(원)', '구매(건)', '매출(원)'])
put(ck, 'A33', '아넬라_프로모_추석세일', border=True)
put(ck, 'B33', f'=SUMIFS({adr("E")},{adr("C")},A33)', f_norm, WON, border=True)
put(ck, 'C33', f'=SUMIFS({adr("H")},{adr("C")},A33)', f_norm, WON, border=True)
put(ck, 'D33', f'=SUMIFS({adr("I")},{adr("C")},A33)', f_norm, WON, border=True)
put(ck, 'A34', '추석세일 종료(9/22) 후에도 10월에 구매가 귀속됨. 월 합계에는 포함, 일별 ROAS는 미사용.', f_note)
put(ck, 'A35', '노출 0인 날에도 전환이 잡히므로, 행사 구간 아넬라 GFA 전환 ROAS(주차별_추이 시트)는 실제보다 높을 가능성.', f_note)
put(ck, 'A36', '5. 기준 확인', f_sec)
notes = ['광고비는 리포트 원본 값. 네이버 VAT 포함 여부는 기존 주간·월간 리포트와 동일 방식으로 처리, 예산표(VAT 미포함)와의 정확한 대응은 확인 필요.',
         '매출은 매체가 보고한 구매(구매완료) 귀속 매출이라 스마트스토어 결제금액과 불일치.',
         '아넬라 GFA 상시 소재는 10/26~10/28 노출이 없는 날에도 전환이 잡혀 귀속 지연으로 판단(사용자 확인).',
         '자사몰 매출 파일의 기간은 10/1~10/28, 판매합계는 실결제 기준(사용자 확인).']
for i, t in enumerate(notes):
    put(ck, f'A{37+i}', f'· {t}', f_norm, al=LEFT); ck.merge_cells(f'A{37+i}:G{37+i}'); ck.row_dimensions[37 + i].height = 28

# =====================================================================
# 주차별_추이
# =====================================================================
wk = sheet('주차별_추이', '2E75B6')
widths(wk, {'A': 30, 'B': 14, 'C': 14, 'D': 14, 'E': 14, 'F': 14, 'G': 16, 'H': 14, 'I': 14})
put(wk, 'A1', '기간별 추이 (상시/행사 구분)', f_title)
put(wk, 'A2', '모든 값은 데이터 시트 SUMIFS 수식. 기간 구분은 가정 시트에서 변경 가능. 광고비에 브랜드검색 정액(일할) 포함.', f_note)
WKR = {}
PC = ['B', 'C', 'D', 'E', 'F']
def wk_block(r0, brand):
    put(wk, f'A{r0}', brand, f_sec)
    put(wk, f'B{r0}', '기간별 광고 성과와 스마트스토어 매출', f_note)
    hdr(wk, r0 + 1, 1, ['지표'], height=22)
    for j, p in enumerate(['P1', 'P2', 'P3', 'P4', 'P5']):
        put(wk, f'{PC[j]}{r0+1}', f"={S_GA}!A{PROW[p]}", f_hdr, fill=fill_hdr, al=CEN, border=True)
    for c, t in zip('GHI', ['10월 누계', '상시 구간', '행사 구간']):
        put(wk, f'{c}{r0+1}', t, f_hdr, fill=fill_hdr, al=CEN, border=True)
    labels = ['기간', '구분', '일수', '광고비(원)', '매출(원)', 'ROAS', '구매(건)', '클릭', '객단가(원)', '일평균 광고비(원)', '일평균 매출(원)',
              '스마트스토어 일반상품 순매출(원)', '(참고) 공동구매 순매출, 집계 제외분(원)', '스스향 광고비(원, 메타 자사몰 제외)', '스마트스토어 통합 ROAS (순매출÷스스향 광고비)']
    for k, t in enumerate(labels):
        put(wk, f'A{r0+2+k}', t, f_norm, border=True, al=LEFT)
    R = {n: r0 + 2 + k for k, n in enumerate(['lab', 'kind', 'days', 'cost', 'rev', 'roas', 'buy', 'clk', 'aov', 'dcost', 'drev', 'ss', 'ssg', 'ssad', 'mer'])}
    WKR[brand] = R
    for j, p in enumerate(['P1', 'P2', 'P3', 'P4', 'P5']):
        c = PC[j]; gr = PROW[p]
        put(wk, f'{c}{R["lab"]}', f'=TEXT({S_GA}!B{gr},"m/d")&"~"&TEXT({S_GA}!C{gr},"m/d")', al=CEN, border=True)
        put(wk, f'{c}{R["kind"]}', f"={S_GA}!E{gr}", al=CEN, border=True)
        put(wk, f'{c}{R["days"]}', f"={S_GA}!D{gr}", fmt=INT, border=True)
        for key, col in [('cost', 'E'), ('rev', 'I'), ('buy', 'H'), ('clk', 'G')]:
            put(wk, f'{c}{R[key]}', f'=SUMIFS({adr(col)},{adr("A")},$A${r0},{adr("D")},{c}${r0+1})', fmt=WON, border=True)
        put(wk, f'{c}{R["ss"]}', f'=SUMIFS({sdr("F")},{sdr("A")},$A${r0},{sdr("C")},"일반",{sdr("B")},">="&{S_GA}!B{gr},{sdr("B")},"<="&{S_GA}!C{gr})', fmt=WON, border=True)
        put(wk, f'{c}{R["ssg"]}', f'=SUMIFS({sdr("F")},{sdr("A")},$A${r0},{sdr("C")},"공구",{sdr("B")},">="&{S_GA}!B{gr},{sdr("B")},"<="&{S_GA}!C{gr})', fmt=WON, border=True)
        put(wk, f'{c}{R["ssad"]}', f'={c}{R["cost"]}-SUMIFS({adr("E")},{adr("A")},$A${r0},{adr("D")},{c}${r0+1},{adr("B")},"메타 자사몰")', fmt=WON, border=True)
    # totals
    put(wk, f'G{R["lab"]}', f'=TEXT({S_GA}!B5,"m/d")&"~"&TEXT({S_GA}!B6,"m/d")', al=CEN, border=True)
    put(wk, f'G{R["kind"]}', '-', al=CEN, border=True)
    put(wk, f'H{R["lab"]}', '상시', al=CEN, border=True); put(wk, f'I{R["lab"]}', '행사', al=CEN, border=True)
    put(wk, f'H{R["kind"]}', '상시', al=CEN, border=True); put(wk, f'I{R["kind"]}', '행사', al=CEN, border=True)
    put(wk, f'G{R["days"]}', f'=SUM(B{R["days"]}:F{R["days"]})', fmt=INT, border=True)
    put(wk, f'H{R["days"]}', f'=SUMIF($B${R["kind"]}:$F${R["kind"]},H${R["kind"]},$B{R["days"]}:$F{R["days"]})', fmt=INT, border=True)
    put(wk, f'I{R["days"]}', f'=SUMIF($B${R["kind"]}:$F${R["kind"]},I${R["kind"]},$B{R["days"]}:$F{R["days"]})', fmt=INT, border=True)
    for key in ['cost', 'rev', 'buy', 'clk', 'ss', 'ssg', 'ssad']:
        put(wk, f'G{R[key]}', f'=SUM(B{R[key]}:F{R[key]})', fmt=WON, border=True)
        for c in 'HI':
            put(wk, f'{c}{R[key]}', f'=SUMIF($B${R["kind"]}:$F${R["kind"]},{c}${R["kind"]},$B{R[key]}:$F{R[key]})', fmt=WON, border=True)
    for c in 'BCDEFGHI':
        put(wk, f'{c}{R["roas"]}', f'=IF({c}{R["cost"]}=0,0,{c}{R["rev"]}/{c}{R["cost"]})', f_bold, PCT, border=True)
        put(wk, f'{c}{R["aov"]}', f'=IF({c}{R["buy"]}=0,0,{c}{R["rev"]}/{c}{R["buy"]})', fmt=WON, border=True)
        put(wk, f'{c}{R["dcost"]}', f'=IF({c}{R["days"]}=0,0,{c}{R["cost"]}/{c}{R["days"]})', fmt=WON, border=True)
        put(wk, f'{c}{R["drev"]}', f'=IF({c}{R["days"]}=0,0,{c}{R["rev"]}/{c}{R["days"]})', fmt=WON, border=True)
        put(wk, f'{c}{R["mer"]}', f'=IF({c}{R["ssad"]}=0,0,{c}{R["ss"]}/{c}{R["ssad"]})', f_bold, PCT, border=True)
    return R
wk_block(4, '아넬라')
wk_block(25, '로지오가닉')

def media_roas(r0, brand_cell_ref, brand):
    put(wk, f'A{r0}', f'{brand} 매체별 ROAS', f_sec)
    hdr(wk, r0 + 1, 1, ['매체'])
    for j, p in enumerate(['P1', 'P2', 'P3', 'P4', 'P5']):
        put(wk, f'{PC[j]}{r0+1}', f"={S_GA}!A{PROW[p]}", f_hdr, fill=fill_hdr, al=CEN, border=True)
    for c, t in zip('GHI', ['10월 누계', '상시 구간', '행사 구간']):
        put(wk, f'{c}{r0+1}', t, f_hdr, fill=fill_hdr, al=CEN, border=True)
    for k, m in enumerate(MEDIA7):
        r = r0 + 2 + k
        put(wk, f'A{r}', m, border=True)
        for c in PC:
            crit = f'{adr("A")},{brand_cell_ref},{adr("B")},$A{r},{adr("D")},{c}${r0+1}'
            put(wk, f'{c}{r}', f'=IF(SUMIFS({adr("E")},{crit})=0,0,SUMIFS({adr("I")},{crit})/SUMIFS({adr("E")},{crit}))', fmt=PCTD, border=True)
        crit = f'{adr("A")},{brand_cell_ref},{adr("B")},$A{r}'
        put(wk, f'G{r}', f'=IF(SUMIFS({adr("E")},{crit})=0,0,SUMIFS({adr("I")},{crit})/SUMIFS({adr("E")},{crit}))', f_bold, PCTD, border=True)
        for c, kind in [('H', '상시'), ('I', '행사')]:
            crit2 = crit + f',{adr("J")},"{kind}"'
            put(wk, f'{c}{r}', f'=IF(SUMIFS({adr("E")},{crit2})=0,0,SUMIFS({adr("I")},{crit2})/SUMIFS({adr("E")},{crit2}))', f_bold, PCTD, border=True)
    put(wk, f'A{r0+10}', '광고비가 0이면 "-"로 표시. 브랜드검색은 정액 대비 매출, GFA 전환 로지오가닉은 10/26 이후만 집행.', f_note)
media_roas(46, '$A$4', '아넬라')
media_roas(58, '$A$25', '로지오가닉')
wk.freeze_panes = 'B4'

# =====================================================================
# 10월_실적
# =====================================================================
rs = sheet('10월_실적', '548235')
widths(rs, {'A': 22, 'B': 13, 'C': 13, 'D': 13, 'E': 11, 'F': 12, 'G': 10, 'H': 9, 'I': 10, 'J': 9, 'K': 9, 'L': 13, 'M': 10, 'N': 11, 'O': 10, 'P': 10})
put(rs, 'A1', '10월 광고 실적 (10/1~10/28) 및 목표 대비 달성률', f_title)
put(rs, 'A2', '금액 단위 만원(별도 표기 제외). 10월 목표는 별도 설정이 없어 4분기 KPI 보고안의 ROAS 하한과 10월 광고예산을 적용한 소급 기준.', f_note)
RSR = {}
def rs_block(r0, brand):
    put(rs, f'A{r0}', brand, f_sec)
    hdr(rs, r0 + 1, 1, ['매체', '10월 월예산', '일할 예산', '광고비', '일할 대비 소진율', '노출', '클릭', 'CTR', 'CPC(원)', '구매(건)', 'CVR', '매출', 'ROAS', 'CPA(원)', '광고비 비중', '매출 비중'], height=32)
    first = r0 + 2
    for k, m in enumerate(MEDIA7):
        r = first + k
        put(rs, f'A{r}', m, border=True)
        if m == '브랜드검색':
            put(rs, f'B{r}', f"={S_GA}!$B$38*{S_GA}!$B$8/10000", fmt=MAN, border=True)
        else:
            put(rs, f'B{r}', f"=SUMIFS({S_GA}!$C$16:$C$27,{S_GA}!$A$16:$A$27,$A${r0},{S_GA}!$B$16:$B$27,$A{r})", fmt=MAN, border=True)
        put(rs, f'C{r}', f"=B{r}*{S_GA}!$B$7/{S_GA}!$B$8", fmt=MAN, border=True)
        crit = f'{adr("A")},$A${r0},{adr("B")},$A{r}'
        put(rs, f'D{r}', f'=SUMIFS({adr("E")},{crit})/10000', fmt=MAN, border=True)
        put(rs, f'F{r}', f'=SUMIFS({adr("F")},{crit})', fmt=INT, border=True)
        put(rs, f'G{r}', f'=SUMIFS({adr("G")},{crit})', fmt=INT, border=True)
        put(rs, f'J{r}', f'=SUMIFS({adr("H")},{crit})', fmt=INT, border=True)
        put(rs, f'L{r}', f'=SUMIFS({adr("I")},{crit})/10000', fmt=MAN, border=True)
    tot = first + len(MEDIA7)
    last = tot - 1
    put(rs, f'A{tot}', f'{brand} 합계', f_bold, fill=fill_sub, border=True)
    for c in 'BCDFGJL':
        put(rs, f'{c}{tot}', f'=SUM({c}{first}:{c}{last})', f_bold, MAN if c in 'BCDL' else INT, fill_sub, border=True)
    for r in list(range(first, tot + 1)):
        bold = f_bold if r == tot else f_norm
        fill = fill_sub if r == tot else None
        put(rs, f'E{r}', f'=IF(C{r}=0,0,D{r}/C{r})', bold, PCTD, fill, border=True)
        put(rs, f'H{r}', f'=IF(F{r}=0,0,G{r}/F{r})', bold, '0.00%', fill, border=True)
        put(rs, f'I{r}', f'=IF(G{r}=0,0,D{r}*10000/G{r})', bold, WON, fill, border=True)
        put(rs, f'K{r}', f'=IF(G{r}=0,0,J{r}/G{r})', bold, '0.0%', fill, border=True)
        put(rs, f'M{r}', f'=IF(D{r}=0,0,L{r}/D{r})', f_bold, PCT, fill, border=True)
        put(rs, f'N{r}', f'=IF(J{r}=0,0,D{r}*10000/J{r})', bold, WON, fill, border=True)
        put(rs, f'O{r}', f'=IF(D${tot}=0,0,D{r}/D${tot})', bold, PCT, fill, border=True)
        put(rs, f'P{r}', f'=IF(L${tot}=0,0,L{r}/L${tot})', bold, PCT, fill, border=True)
    RSR[brand] = dict(first=first, tot=tot, title=r0)
rs_block(4, '아넬라')
rs_block(16, '로지오가닉')
put(rs, 'A27', '브랜드검색은 월 정액을 일할 안분한 값이며 예산표(네이버 VAT 미포함)에는 미포함. 합계 광고비·ROAS는 브랜드검색 정액을 포함한 값으로 주간·월간 리포트와 동일 기준.', f_note)

# --- 달성 table 1
a_t, r_t = RSR['아넬라']['tot'], RSR['로지오가닉']['tot']
put(rs, 'A29', '① 10월 목표(소급) 대비 28일 누계 실적', f_sec)
hdr(rs, 30, 1, ['브랜드', '월 광고비 계획', 'KPI ROAS 하한', '목표 매출 (계획×하한)', '28일 광고비', '28일 매출', '28일 ROAS', 'ROAS 달성률', '판정'], height=32)
cols_gh = {'아넬라': 'B', '로지오가닉': 'C'}
for i, b in enumerate(BRANDS):
    r = 31 + i; t = RSR[b]['tot']
    put(rs, f'A{r}', b, f_bold, border=True)
    put(rs, f'B{r}', f'=B{t}', fmt=MAN, border=True)
    put(rs, f'C{r}', f"={S_GA}!{cols_gh[b]}42", fmt=PCT, border=True)
    put(rs, f'D{r}', f'=B{r}*C{r}', fmt=MAN, border=True)
    put(rs, f'E{r}', f'=D{t}', fmt=MAN, border=True)
    put(rs, f'F{r}', f'=L{t}', fmt=MAN, border=True)
    put(rs, f'G{r}', f'=IF(E{r}=0,0,F{r}/E{r})', f_bold, PCT, border=True)
    put(rs, f'H{r}', f'=IF(C{r}=0,0,G{r}/C{r})', f_bold, PCT, border=True)
    put(rs, f'I{r}', f'=IF(H{r}>=1,"달성",IF(H{r}>={S_GA}!$B$12,"근접","미달"))', f_bold, al=CEN, border=True)
put(rs, 'A33', '합계', f_bold, fill=fill_sub, border=True)
for c in 'BDEF':
    put(rs, f'{c}33', f'={c}31+{c}32', f_bold, MAN, fill_sub, border=True)
put(rs, 'C33', '=IF(B33=0,0,D33/B33)', f_bold, PCT, fill_sub, border=True)
put(rs, 'G33', '=IF(E33=0,0,F33/E33)', f_bold, PCT, fill_sub, border=True)
put(rs, 'H33', '=IF(C33=0,0,G33/C33)', f_bold, PCT, fill_sub, border=True)
put(rs, 'I33', f'=IF(H33>=1,"달성",IF(H33>={S_GA}!$B$12,"근접","미달"))', f_bold, al=CEN, fill=fill_sub, border=True)
put(rs, 'A34', '월 광고비 계획 = 10월 예산 합계 + 브랜드검색 정액(31일분). 합계 행의 ROAS 하한은 브랜드별 목표 매출을 계획 광고비로 가중한 값.', f_note)

# --- table 2 month-end projection
put(rs, 'A36', '② 월말 예상 (10/29~31은 10/26~28 일평균이 이어진다고 가정)', f_sec)
hdr(rs, 37, 1, ['브랜드', '월말 예상 광고비', '월말 예상 매출', '예상 ROAS', '예상 ROAS 달성률', '예상 매출 ÷ 목표 매출', '예산 소진율 (예상 광고비÷계획)', '월말 예상 구매(건)', 'KPI 광고 기준 목표 주문(건)', '주문 달성률', '판정 (예상 ROAS 기준)'], height=44)
for i, b in enumerate(BRANDS):
    r = 38 + i; t = RSR[b]['tot']; w = WKR[b]; t1 = 31 + i
    put(rs, f'A{r}', b, f_bold, border=True)
    put(rs, f'B{r}', f"=E{t1}+{S_WK}!F{w['dcost']}*{S_GA}!$B$9/10000", fmt=MAN, border=True)
    put(rs, f'C{r}', f"=F{t1}+{S_WK}!F{w['drev']}*{S_GA}!$B$9/10000", fmt=MAN, border=True)
    put(rs, f'D{r}', f'=IF(B{r}=0,0,C{r}/B{r})', f_bold, PCT, border=True)
    put(rs, f'E{r}', f'=IF(C{t1}=0,0,D{r}/C{t1})', f_bold, PCT, border=True)
    put(rs, f'F{r}', f'=IF(D{t1}=0,0,C{r}/D{t1})', fmt=PCT, border=True)
    put(rs, f'G{r}', f'=IF(B{t1}=0,0,B{r}/B{t1})', fmt=PCT, border=True)
    put(rs, f'H{r}', f"=J{t}+{S_WK}!F{w['buy']}/{S_WK}!F{w['days']}*{S_GA}!$B$9", fmt=INT, border=True)
    put(rs, f'I{r}', f"={S_GA}!{cols_gh[b]}43", fmt=INT, border=True)
    put(rs, f'J{r}', f'=IF(I{r}=0,0,H{r}/I{r})', fmt=PCT, border=True)
    put(rs, f'K{r}', f'=IF(E{r}>=1,"달성",IF(E{r}>={S_GA}!$B$12,"근접","미달"))', f_bold, al=CEN, border=True)
put(rs, 'A40', '합계', f_bold, fill=fill_sub, border=True)
for c in 'BC':
    put(rs, f'{c}40', f'={c}38+{c}39', f_bold, MAN, fill_sub, border=True)
put(rs, 'D40', '=IF(B40=0,0,C40/B40)', f_bold, PCT, fill_sub, border=True)
put(rs, 'E40', '=IF(C33=0,0,D40/C33)', f_bold, PCT, fill_sub, border=True)
put(rs, 'F40', '=IF(D33=0,0,C40/D33)', f_bold, PCT, fill_sub, border=True)
put(rs, 'G40', '=IF(B33=0,0,B40/B33)', f_bold, PCT, fill_sub, border=True)
put(rs, 'H40', '=H38+H39', f_bold, INT, fill_sub, border=True)
put(rs, 'I40', '=I38+I39', f_bold, INT, fill_sub, border=True)
put(rs, 'J40', '=IF(I40=0,0,H40/I40)', f_bold, PCT, fill_sub, border=True)
put(rs, 'K40', f'=IF(E40>=1,"달성",IF(E40>={S_GA}!$B$12,"근접","미달"))', f_bold, al=CEN, fill=fill_sub, border=True)
put(rs, 'A41', '월말 예상 광고비·매출 = 28일 누계 + 10/26~28 일평균 × 잔여일수. 10/26~28이 넾다세일 구간이라 상시 구간 대비 높게 산출될 가능성. KPI 목표 주문은 4분기 KPI 보고안의 월 환산값.', f_note)

# --- table 3 anomaly sensitivity
put(rs, 'A43', '③ 로지오가닉 메타 스스 이상치(10/24~25) 제외 시 (점검 시트 참조)', f_sec)
hdr(rs, 44, 1, ['구분', '28일 광고비', '28일 매출', '28일 ROAS', 'KPI 하한 대비'], height=24)
put(rs, 'A45', '원본', border=True)
put(rs, 'B45', '=E32', fmt=MAN, border=True); put(rs, 'C45', '=F32', fmt=MAN, border=True)
put(rs, 'A46', '보정 (이상치 이틀 제외)', border=True)
put(rs, 'B46', f"=E32-{S_CK}!D14/10000", fmt=MAN, border=True); put(rs, 'C46', f"=F32-{S_CK}!E14/10000", fmt=MAN, border=True)
for r in (45, 46):
    put(rs, f'D{r}', f'=IF(B{r}=0,0,C{r}/B{r})', f_bold, PCT, border=True)
    put(rs, f'E{r}', f'=IF(C32=0,0,D{r}/C32)', f_bold, PCT, border=True)
rs.freeze_panes = 'B4'
for rng in ['I31:I33', 'K38:K40']:
    pass
for rng, first in [('I31:I33', 'I31'), ('K38:K40', 'K38')]:
    rs.conditional_formatting.add(rng, FormulaRule(formula=[f'{first}="달성"'], fill=PatternFill('solid', bgColor='C6E0B4', fgColor='C6E0B4')))
    rs.conditional_formatting.add(rng, FormulaRule(formula=[f'{first}="근접"'], fill=PatternFill('solid', bgColor='FFE699', fgColor='FFE699')))
    rs.conditional_formatting.add(rng, FormulaRule(formula=[f'{first}="미달"'], fill=PatternFill('solid', bgColor='F8CBAD', fgColor='F8CBAD')))

# =====================================================================
# 목표_11-12월
# =====================================================================
tg = sheet('목표_11-12월', 'C55A11')
widths(tg, {'A': 22, 'B': 16, 'C': 13, 'D': 13, 'E': 13, 'F': 13, 'G': 12, 'H': 12, 'I': 11, 'J': 11, 'K': 12, 'L': 11, 'M': 11, 'N': 11, 'O': 13, 'P': 13, 'Q': 13})
put(tg, 'A1', '11~12월 목표 (보수 / 기준 / 도전)', f_title)
put(tg, 'A2', '10월 실적의 상시 ROAS·행사 ROAS를 매체별로 산출한 뒤 월별 행사일수에 따라 혼합해 목표 ROAS와 목표 매출 계산. 노란 칸 가정 수정 시 전체 재계산.', f_note)
put(tg, 'A4', '시나리오 가정 (입력)', f_sec)
hdr(tg, 5, 1, ['항목', '보수', '기준', '도전', '근거'])
put(tg, 'A6', '11월 행사일수 (일)', border=True)
for c, v in zip('BCD', [8, 11, 14]):
    put(tg, f'{c}6', v, f_in, INT, fill_in, RIGHT, True)
put(tg, 'E6', '보수 = 넾다세일 11/1~11/8 확정분 8일. 기준 = 여기에 아넬라 오늘끝딜(11월 2주 선정 대기) 등 3일 가정. 도전 = 영업팀 블랙프라이데이 행사 3일을 더 가정. 기준·도전의 추가 일수는 가정이며 일정 확정 시 수정', f_note, al=LEFT)
tg.merge_cells('E6:Q6'); tg.row_dimensions[6].height = 30
put(tg, 'A7', '12월 행사일수 (일)', border=True)
for c, v in zip('BCD', [5, 8, 11]):
    put(tg, f'{c}7', v, f_in, INT, fill_in, RIGHT, True)
put(tg, 'E7', '12월 일정은 미정. 10월 행사 비중(28일 중 10일, 약 36%)을 31일로 환산한 11일을 도전으로, 기준·보수는 그보다 낮게 가정', f_note, al=LEFT)
tg.merge_cells('E7:Q7')
put(tg, 'A8', '이상치 보정 적용 (1=적용, 0=미적용)', border=True)
put(tg, 'B8', 1, f_in, INT, fill_in, RIGHT, True)
put(tg, 'E8', '1이면 로지오가닉 메타 스스의 10/24~25를 상시 ROAS 계산에서 제외 (점검 시트 1번)', f_note)
tg.merge_cells('E8:Q8')

TGR = {}
def tg_month(r0, month, gcol, dayscell, eventrow):
    put(tg, f'A{r0}', f'{month} 매체별 목표', f_sec)
    hdr(tg, r0 + 1, 1, ['브랜드', '매체', '10월 상시 광고비(원)', '10월 상시 매출(원)', '10월 행사 광고비(원)', '10월 행사 매출(원)', '보정 광고비(원)', '보정 매출(원)',
                        '상시 ROAS', '행사 ROAS', f'{month} 광고비(만원)', '목표 ROAS 보수', '목표 ROAS 기준', '목표 ROAS 도전', '목표 매출 보수(만원)', '목표 매출 기준(만원)', '목표 매출 도전(만원)'], height=44)
    r = r0 + 2
    rows_by_brand = {}
    for b in BRANDS:
        start = r
        for m in MEDIA7:
            put(tg, f'A{r}', b, border=True); put(tg, f'B{r}', m, border=True)
            base = f'{adr("A")},$A{r},{adr("B")},$B{r}'
            put(tg, f'C{r}', f'=SUMIFS({adr("E")},{base},{adr("J")},"상시")', fmt=WON, border=True)
            put(tg, f'D{r}', f'=SUMIFS({adr("I")},{base},{adr("J")},"상시")', fmt=WON, border=True)
            put(tg, f'E{r}', f'=SUMIFS({adr("E")},{base},{adr("J")},"행사")', fmt=WON, border=True)
            put(tg, f'F{r}', f'=SUMIFS({adr("I")},{base},{adr("J")},"행사")', fmt=WON, border=True)
            put(tg, f'G{r}', f"=$B$8*SUMIFS({S_CK}!$D$14:$D$14,{S_CK}!$A$14:$A$14,$A{r},{S_CK}!$B$14:$B$14,$B{r})", fmt=WON, border=True)
            put(tg, f'H{r}', f"=$B$8*SUMIFS({S_CK}!$E$14:$E$14,{S_CK}!$A$14:$A$14,$A{r},{S_CK}!$B$14:$B$14,$B{r})", fmt=WON, border=True)
            # base ROAS with fallbacks
            put(tg, f'I{r}', f'=IF(C{r}-G{r}>0,(D{r}-H{r})/(C{r}-G{r}),IF(E{r}>0,F{r}/E{r},0))', f_bold, PCTD, border=True)
            put(tg, f'J{r}', f'=IF(E{r}>0,F{r}/E{r},I{r})', f_bold, PCTD, border=True)
            if m == '브랜드검색':
                put(tg, f'K{r}', f"={S_GA}!$B$38*{dayscell}/10000", fmt=MAN, border=True)
            else:
                put(tg, f'K{r}', f"=SUMIFS({S_GA}!${gcol}$16:${gcol}$27,{S_GA}!$A$16:$A$27,$A{r},{S_GA}!$B$16:$B$27,$B{r})", fmt=MAN, border=True)
            for j, sc in enumerate('BCD'):
                put(tg, f'{"LMN"[j]}{r}', f'=(${sc}${eventrow}*$J{r}+({dayscell}-${sc}${eventrow})*$I{r})/{dayscell}', f_bold, PCTD, border=True)
                put(tg, f'{"OPQ"[j]}{r}', f'=$K{r}*{"LMN"[j]}{r}', fmt=MAN, border=True)
            r += 1
        end = r - 1
        put(tg, f'A{r}', b, f_bold, fill=fill_sub, border=True); put(tg, f'B{r}', '합계', f_bold, fill=fill_sub, border=True)
        for c in 'CDEFGH':
            put(tg, f'{c}{r}', '', fill=fill_sub, border=True)
        put(tg, f'I{r}', '', fill=fill_sub, border=True); put(tg, f'J{r}', '', fill=fill_sub, border=True)
        put(tg, f'K{r}', f'=SUM(K{start}:K{end})', f_bold, MAN, fill_sub, border=True)
        for j in range(3):
            o = "OPQ"[j]
            put(tg, f'{o}{r}', f'=SUM({o}{start}:{o}{end})', f_bold, MAN, fill_sub, border=True)
            put(tg, f'{"LMN"[j]}{r}', f'=IF($K{r}=0,0,{o}{r}/$K{r})', f_bold, PCT, fill_sub, border=True)
        rows_by_brand[b] = (start, end, r)
        r += 1
    # judgement rows
    put(tg, f'A{r+1}', 'KPI ROAS 하한 대비 판정', f_bold)
    jr = {}
    for i, b in enumerate(BRANDS):
        rr = r + 2 + i
        tot = rows_by_brand[b][2]
        put(tg, f'A{rr}', b, border=True)
        put(tg, f'B{rr}', f"={S_GA}!{cols_gh[b]}42", fmt=PCT, border=True)
        put(tg, f'C{rr}', '하한', f_note)
        for j in range(3):
            put(tg, f'{"LMN"[j]}{rr}', f'=IF({"LMN"[j]}{tot}>=$B{rr},"달성","미달")', f_bold, al=CEN, border=True)
        # required event days
        s, e, _ = rows_by_brand[b]
        put(tg, f'O{rr}', f'=IF(SUMPRODUCT(K{s}:K{e},J{s}:J{e}-I{s}:I{e})=0,0,($B{rr}*K{tot}-SUMPRODUCT(K{s}:K{e},I{s}:I{e}))*{dayscell}/SUMPRODUCT(K{s}:K{e},J{s}:J{e}-I{s}:I{e}))', f_bold, '0.0"일"', border=True)
        put(tg, f'P{rr}', '← 하한 달성에 필요한 행사일수(현재 예산 기준)', f_note)
        jr[b] = rr
    hr = r + 1
    put(tg, f'L{hr}', '보수', f_hdr, fill=fill_hdr, al=CEN); put(tg, f'M{hr}', '기준', f_hdr, fill=fill_hdr, al=CEN); put(tg, f'N{hr}', '도전', f_hdr, fill=fill_hdr, al=CEN)
    put(tg, f'O{hr}', '필요 행사일수', f_hdr, fill=fill_hdr, al=CEN)
    TGR[month] = dict(rows=rows_by_brand, judge=jr, end=r + 3)
    rng = f'L{r+2}:N{r+3}'
    tg.conditional_formatting.add(rng, FormulaRule(formula=[f'L{r+2}="달성"'], fill=PatternFill('solid', bgColor='C6E0B4', fgColor='C6E0B4')))
    tg.conditional_formatting.add(rng, FormulaRule(formula=[f'L{r+2}="미달"'], fill=PatternFill('solid', bgColor='F8CBAD', fgColor='F8CBAD')))
    return r + 3
e1 = tg_month(10, '11월', 'D', f"{S_GA}!$B$10", 6)
e2 = tg_month(e1 + 3, '12월', 'E', f"{S_GA}!$B$11", 7)

# SS sales target
r0 = e2 + 3
put(tg, f'A{r0}', '스마트스토어 일반상품 순매출 목표 (공동구매 제외, 만원)', f_sec)
hdr(tg, r0 + 1, 1, ['브랜드', '월', '스스향 광고비(만원)', '10월 통합 ROAS 상시', '10월 통합 ROAS 행사', '혼합 보수', '혼합 기준', '혼합 도전', '목표 순매출 보수', '목표 순매출 기준', '목표 순매출 도전'], height=44)
ssr = {}
r = r0 + 2
for b in BRANDS:
    w = WKR[b]
    for month, dayscell, ev in [('11월', f"{S_GA}!$B$10", 6), ('12월', f"{S_GA}!$B$11", 7)]:
        t = TGR[month]['rows'][b]
        s, e, tot = t
        put(tg, f'A{r}', b, border=True); put(tg, f'B{r}', month, border=True)
        jam = [x for x in range(s, e + 1) if True][-1]  # last media row is 메타 자사몰
        put(tg, f'C{r}', f'=K{tot}-K{jam}', fmt=MAN, border=True)
        put(tg, f'D{r}', f"={S_WK}!H{w['mer']}", fmt=PCT, border=True)
        put(tg, f'E{r}', f"={S_WK}!I{w['mer']}", fmt=PCT, border=True)
        for j, sc in enumerate('BCD'):
            put(tg, f'{"FGH"[j]}{r}', f'=(${sc}${ev}*$E{r}+({dayscell}-${sc}${ev})*$D{r})/{dayscell}', f_bold, PCT, border=True)
            put(tg, f'{"IJK"[j]}{r}', f'=$C{r}*{"FGH"[j]}{r}', f_bold, MAN, border=True)
        ssr[(b, month)] = r
        r += 1
put(tg, f'A{r}', '통합 ROAS = 스마트스토어 일반상품 순매출 ÷ 스스향 광고비(메타 자사몰 제외). 로지오가닉은 상시 구간에 이상치 이틀이 그대로 포함되어 목표 순매출이 다소 높게 산출될 가능성.', f_note)
put(tg, f'A{r+1}', '목표 매출·순매출은 예산 100% 사용 가정 값. 10월 예산 소진율은 아넬라 약 84%, 로지오가닉 약 80%(10월_실적 시트).', f_note)
put(tg, f'A{r+2}', '아넬라 GFA 전환의 행사 ROAS는 귀속 지연·뷰스루 영향으로 높게 산출되었을 가능성(점검 시트 4번). 해당 매체 목표 매출은 보수적 접근 권장.', f_note)
tg.freeze_panes = 'C4'

# =====================================================================
# 성과분석
# =====================================================================
an = sheet('성과분석', '7030A0')
widths(an, {'A': 12, 'B': 12, 'C': 26, 'D': 34, 'E': 13, 'F': 13, 'G': 58, 'H': 52})
put(an, 'A1', '성과 분석: 잘한 점·이유, 미달·보완', f_title)
put(an, 'A2', '값 열은 다른 시트 수식. "원인"은 주간 리포트와 사용자 설명에서 확인된 운영 이력에 근거하며, 데이터만으로 단정하기 어려운 항목은 "추정"으로 표시.', f_note)
hdr(an, 4, 1, ['구분', '브랜드', '항목', '근거 지표', '값', '비교 기준', '성과 이유 / 원인', '보완 방향'], height=26)
wa, wr_ = WKR['아넬라'], WKR['로지오가닉']
ra, rr_ = RSR['아넬라'], RSR['로지오가닉']
def rsm(brand, media, col):
    R = RSR[brand]; return f"{S_RS}!{col}{R['first'] + MEDIA7.index(media)}"
def wkm(brand, media, col):
    base = 48 if brand == '아넬라' else 60
    return f"{S_WK}!{col}{base + MEDIA7.index(media)}"
items = [
 # (구분, 브랜드, 항목, 근거지표, 값수식, 값포맷, 비교수식, 비교포맷, 이유, 보완)
 ('잘한 점', '아넬라', '파워링크 노출 시간 제한 (10/14~)', 'ROAS 제한 후 (비교: 제한 전, 두 기간 모두 상시 구간)', "='시간제한_검증'!C13", PCT, "='시간제한_검증'!B13", PCT,
  '노출 제외 시간(0~3시, 17~24시) 설정 후 광고비·클릭 약 58% 감소, 구매는 15건에서 14건으로 거의 유지, 전환율 6.1%에서 13.7%로 상승. 행사 없는 두 주 비교라 행사 영향 없음.',
  '매출 볼륨(-22%)이 줄었으므로 추가 모니터링. 같은 기간에 10/21 저효율 키워드 감액이 섞여 있어 시간 제한만의 효과로 단정하지 않기.'),
 ('잘한 점', '아넬라', '저효율 키워드·지면 정리로 예산 효율화', '28일 ROAS (비교: KPI 하한)', f"={S_RS}!G31", PCT, f"={S_RS}!C31", PCT,
  '저효율 키워드·지면 OFF/감액과 노출 시간 제한으로 일할 예산의 약 84%만 집행, ROAS는 하한에 근접. 해당 ROAS에 행사 구간 영향이 포함돼 있어 마케팅 성과는 상시 구간 지표(미달·보완 항목)로 별도 판단.',
  '정리 효과는 11/8 넾다세일 종료 후 상시 구간 데이터로 검증.'),
 ('잘한 점', '아넬라', '상시 구간에서 쇼핑검색이 핵심 매체', '쇼핑검색 상시 구간 ROAS (비교: 파워링크 상시)', f"={wkm('아넬라','쇼핑검색','H')}", PCT, f"={wkm('아넬라','파워링크','H')}", PCT,
  '행사 없는 구간에서도 쇼핑검색이 광고비 비중 최대 매체(약 32%)이며 파워링크보다 효율 높음. 퓨레 캠페인이 매출 견인.',
  '쇼핑검색은 10/26부터 24시간 균등 노출로 원복했으므로 11/8 이후 상시 구간에서 제한 구간(10/14~25)과 비교해 효과를 판단.'),
 ('잘한 점', '아넬라', '메타 시크릿링크가 자사몰 매출의 주축', '시크릿특가 상품 매출 비중 (비교: 시크릿특가 실매출 ROAS)', "='자사몰_매출'!H7", PCT, "='자사몰_매출'!G15", PCT,
  '자사몰 상품명에 "시크릿 특가"가 붙은 상품은 메타 시크릿링크 광고 유입 상품. 해당 상품이 자사몰 매출의 대부분을 차지하며 시크릿링크 광고가 자사몰 매출의 핵심 경로.',
  '시크릿링크 광고 보고 매출이 실제 시크릿특가 매출보다 큼(자사몰_매출 시트). 실매출 기준 ROAS로 판단.'),
 ('잘한 점', '로지오가닉', '메타 시크릿링크가 자사몰 매출의 주축', '시크릿특가 상품 매출 비중 (비교: 시크릿특가 실매출 ROAS)', "='자사몰_매출'!H8", PCT, "='자사몰_매출'!G16", PCT,
  '로지 자사몰도 "시크릿 특가" 상품(쌀과자 퍼프 스틱 떡뻥 요거팝)이 매출 1위. 첫구매 전용 상품도 별도 비중 존재(자사몰_매출 시트).',
  '실매출 ROAS가 KPI 하한(420%)을 넘는지 계속 확인. 개편 썸네일 소재와 병행 테스트.'),
 ('잘한 점', '로지오가닉', 'KPI 하한 상회 (행사 영향이 작은 브랜드)', '28일 ROAS (비교: KPI 하한)', f"={S_RS}!G32", PCT, f"={S_RS}!C32", PCT,
  '로지는 행사·상시 구간 ROAS 차이가 작아(주차별_추이 시트: 상시 약 566%, 행사 약 687%) 행사 의존도 낮음. 메타 스스(과자 상시)와 메타 자사몰(과자 시크릿링크)이 광고비의 절반 이상을 차지하며 평균 ROAS 상승 견인.',
  '메타 스스 이상치 이틀 제외 시에도 하한을 넘는지 오른쪽 항목 참고.'),
 ('잘한 점', '로지오가닉', '이상치 제외 후에도 하한 상회', '보정 ROAS (비교: KPI 하한)', f"={S_RS}!D46", PCT, f"={S_RS}!C32", PCT,
  '메타 스스 10/24~25 제외 시에도 하한 상회. 성과 달성 판단은 이상치와 무관하게 유지.',
  '이상치 원인 확인 후 확정 수치로 교체.'),
 ('잘한 점', '로지오가닉', '메타 자사몰 과자 시크릿링크의 안정성', '28일 ROAS (비교: KPI 하한)', f"={rsm('로지오가닉','메타 자사몰','M')}", PCT, f"={S_RS}!C32", PCT,
  '10/1~4부터 10/26~28까지 5개 구간 모두 430~750% 범위 유지(주차별 ROAS는 주차별_추이 시트). 시크릿링크 구성이 일관되게 구매 창출.',
  '개편된 시크릿링크 썸네일로 소재를 교체할 때 기존 소재와 병행 테스트.'),
 ('미달·보완', '아넬라', 'KPI ROAS 하한 소폭 미달', '28일 ROAS 달성률', f"={S_RS}!H31", PCT, 1, PCT,
  '상시 구간 ROAS가 행사 구간의 약 1/3 수준이라 행사 없는 주에는 하한 크게 하회(주차별_추이 시트).',
  '상시 구간 효율 개선: 저효율 지면·키워드 정리 유지, 소재 교체(차주 2종), 행사 일정 확대 협의.'),
 ('미달·보완', '아넬라', '상시 구간 ROAS', '상시 구간 ROAS (비교: KPI 하한)', f"={S_WK}!H{wa['roas']}", PCT, f"={S_RS}!C31", PCT,
  '10/1~4, 10/5~11, 10/19~25의 합산 ROAS. 4주차(10/19~25)는 전 매체 ROAS 전주 대비 하락, 전환율 16.8%에서 7.0%로 하락(주간 리포트).',
  '쇼핑검색은 24시간 균등노출로 원복 완료(10/26~). 파워링크는 시간 제한 유지, 무전환 지면 제외 후 효율 확인.'),
 ('미달·보완', '아넬라', '예산 소진 부족', '일할 예산 대비 소진율 (10월 28일)', f"={S_RS}!E{ra['tot']}", PCT, 1, PCT,
  '저효율 키워드·지면 OFF/감액, 웰컴박스·기획전 보류로 의도적 감액 영향이 큼. GFA는 10/26부터 상시 소재 노출이 거의 없고 넾다세일 소재 중심 운영으로 추정.',
  '남은 예산을 행사 구간 소재에 집중할지, 효율 우선으로 줄일지 11월 계획에서 결정.'),
 ('미달·보완', '아넬라', '메타 자사몰 효율 하락', '10월 28일 ROAS (비교: 아넬라 메타 자사몰 P5 ROAS)', f"={rsm('아넬라','메타 자사몰','M')}", PCT, f"={wkm('아넬라','메타 자사몰','F')}", PCT,
  '웰컴박스는 신청하기 목적 캠페인이라 구매 ROAS 거의 없음(광고비 약 20만원, 구매 2건), 시크릿링크도 구간 경과에 따라 ROAS 하락. 광고 보고 ROAS와 자사몰 실매출 기준 ROAS 차이는 자사몰_매출 시트 참고.',
  '웰컴박스는 보류 유지. 시크릿링크 개편 썸네일 소재로 교체 후 재측정.'),
 ('미달·보완', '아넬라', '메타 카탈로그(스스)', '캠페인 ROAS', f"=SUMIFS({adr('I')},{adr('C')},\"SUS_아넬라_카탈로그 판매 캠페인_260629\")/SUMIFS({adr('E')},{adr('C')},\"SUS_아넬라_카탈로그 판매 캠페인_260629\")", PCT, 2, PCT,
  '월 ROAS는 양호하나 4주차 189%까지 하락(주간 리포트). 비교 기준 200%는 주간 리포트의 개선 기준.',
  '카탈로그 상품 세트·소재를 영업팀 주력 상품 중심으로 재구성.'),
 ('미달·보완', '아넬라', '쇼핑검색 과자(떡뻥) 캠페인', '캠페인 ROAS (비교: KPI 하한)', f"=SUMIFS({adr('I')},{adr('C')},\"SSA_아넬라_과자(떡뻥) 캠페인\")/SUMIFS({adr('E')},{adr('C')},\"SSA_아넬라_과자(떡뻥) 캠페인\")", PCT, f"={S_RS}!C31", PCT,
  '9월 월간 리포트에서도 과자(떡뻥)_일반K 효율 낮음(ROAS 201%). 대표 단어 경쟁이 강해 증액 시 효율 희석.',
  '롱테일 검색어(퍼프·롱떡뻥) 중심으로 소액 운영, 저효율 검색어 제외.'),
 ('미달·보완', '아넬라', 'GFA 신규 소구 5종 테스트', '집행 구간 ROAS (10/14~16)', f"=SUMIFS({adr('I')},{adr('C')},\"아넬라_신규타겟_소구5종\")/SUMIFS({adr('E')},{adr('C')},\"아넬라_신규타겟_소구5종\")", PCT, 1, PCT,
  '3일 집행에서 구매 1건으로 효율 미확보, 소재 검토를 위해 HOLD 중.',
  '소재 검토 후 재개. 재개 전 소구점·타겟 가설을 팀과 합의.'),
 ('미달·보완', '로지오가닉', '파워링크·쇼핑검색 효율', '파워링크+쇼핑검색 합산 ROAS (브랜드검색 제외, 비교: KPI 하한)',
  f"=(SUMIFS({adr('I')},{adr('A')},\"로지오가닉\",{adr('B')},\"파워링크\")+SUMIFS({adr('I')},{adr('A')},\"로지오가닉\",{adr('B')},\"쇼핑검색\"))/(SUMIFS({adr('E')},{adr('A')},\"로지오가닉\",{adr('B')},\"파워링크\")+SUMIFS({adr('E')},{adr('A')},\"로지오가닉\",{adr('B')},\"쇼핑검색\"))", PCT, f"={S_RS}!C32", PCT,
  '파워링크 대표K·구매K 그룹의 구매 전환 감소 영향이 큼(4주차 주간 리포트). 브랜드검색 포함 시 검색 전체 ROAS 상승.',
  '10/21 파워링크 저효율 키워드 감액 효과를 11/8 이후 상시 구간으로 검증. 쇼핑검색 무전환 지면 제외(10/26) 효과 확인.'),
 ('미달·보완', '로지오가닉', 'GFA 웹전환(넾다세일 한정)', '28일 ROAS (10/26~28 집행)', f"={rsm('로지오가닉','GFA 전환','M')}", PCT, 1, PCT,
  '3일 집행, 구매 6건으로 표본은 작으나 현재 광고비 미회수.',
  '넾다세일 종료(11/8)까지 모니터링 후 지속 여부를 판단하고, 효율이 나오지 않으면 20만원을 메타로 이동.'),
 ('데이터 점검', '로지오가닉', '메타 스스 10/24~25 이상치', '이틀 메타 전환값 ÷ 스토어 결제금액 (배)', f"={S_CK}!D15", '0.0"배"', 1, '0.0"배"',
  '같은 이틀 로지 스토어 전체 결제(공구 포함)보다 메타 전환값이 큼. 실주문이 아닌 집계 이슈 가능성.',
  '메타·네이버 협력광고 담당에 확인. 확인 전까지 목표 산정에서는 제외(점검 시트).'),
 ('데이터 점검', '로지오가닉', '공동구매 의존도', '스토어 결제 총액 중 공동구매 비중', f"={S_CK}!F22", PCT, 0, PCT,
  '공동구매 상품(9차 등) 제외 시 스토어 일반상품 매출은 절반 수준. 광고 귀속 매출은 일반상품 순매출에 근접.',
  '일반상품 기준으로 광고 효율을 판단하되, 공동구매 기간의 유입·전환이 광고 성과에 섞이는지 영업팀과 확인.'),
]
for i, (sec, b, item, metric, vf, vfmt, cf, cfmt, why, fix) in enumerate(items):
    r = 5 + i
    put(an, f'A{r}', sec, f_bold, al=CEN, border=True)
    put(an, f'B{r}', b, al=CEN, border=True)
    put(an, f'C{r}', item, f_bold, al=LEFT, border=True)
    put(an, f'D{r}', metric, f_note, al=LEFT, border=True)
    put(an, f'E{r}', vf, f_bold, vfmt, al=RIGHT, border=True)
    put(an, f'F{r}', cf, f_norm, cfmt, al=RIGHT, border=True)
    put(an, f'G{r}', why, al=LEFT, border=True)
    put(an, f'H{r}', fix, al=LEFT, border=True)
    an.row_dimensions[r].height = 58
    fillc = {'잘한 점': 'E2EFDA', '미달·보완': 'FCE4D6', '데이터 점검': 'FFF2CC'}[sec]
    an[f'A{r}'].fill = PatternFill('solid', fgColor=fillc)
an.freeze_panes = 'D5'
AN_LAST = 4 + len(items)
ANR = {it[2]: 5 + i for i, it in enumerate(items)}

# =====================================================================
# 실행계획
# =====================================================================
ep = sheet('실행계획', 'BF8F00')
widths(ep, {'A': 6, 'B': 12, 'C': 16, 'D': 14, 'E': 52, 'F': 44, 'G': 34, 'H': 14})
put(ep, 'A1', '11~12월 실행계획', f_title)
put(ep, 'A2', '근거 열은 10월 데이터 또는 사용자 전달 사항. 상태가 "제안"인 항목은 담당·일정 확인 필요.', f_note)
hdr(ep, 4, 1, ['#', '브랜드', '매체·영역', '시기', '실행 내용', '목적 / 기대 효과', '근거 · 선행 조건', '상태'], height=26)
plan = [
 ('공통', '메타 전체', '11/1~11/8', '넾다세일 기간 메타는 인스타그램 지면만 고정으로 운영하고, 저효율 지면은 OFF/감액을 유지. 광고 상품을 상시 상품(아넬라 퓨레 6+6, 로지 쌀과자 8봉)에서 영업팀 주력 상품으로 전환.', '행사 구간 전환 극대화, 저효율 지면 비용 제거', '넾다세일 11/8 종료. 10월 중순부터 저효율 지면 OFF/감액 중', '진행 중'),
 ('공통', '네이버 검색·GFA', '11/9~', '넾다세일 종료 후 상시 구간 데이터로 10/14(시간대 제한), 10/21(파워링크 저효율 키워드 감액), 10/26(무전환 지면 제외) 조치의 효과를 비교하고 유지·원복을 결정.', '조치 효과 검증, 상시 ROAS 개선', '상시 구간 ROAS가 하한 달성의 관건 (성과분석 시트)', '예정'),
 ('공통', '영업팀 협의', '11월 1주 내', '블랙프라이데이 등 11월 행사 일정을 영업팀과 협의. 행사 일정이 확정되면 목표 시트의 행사일수를 수정해 목표와 예산을 재산정.', '행사일수 확보 (필요 행사일수는 목표 시트 참고)', '행사 일정은 영업팀 소관, 현재 확정 일정은 넾다세일뿐', '제안'),
 ('공통', '메타 파트너십', '넾다세일 이후', '자사몰 첫구매 혜택 영상 소재를 파트너십광고로 운영.', '자사몰 신규 구매 확보, 행사 외 시기 효율 방어', '자사몰 첫구매 전용 상품 매출 규모는 자사몰_매출 시트 참고(로지 쪽이 큼). 개편 시크릿링크 썸네일 소재와 병행 테스트', '예정'),
 ('아넬라', '네이버 오늘끝딜', '11월 2주', '오늘끝딜 선정 시 해당 일자 검색·메타 예산을 증액하고 소재를 행사 소구로 교체. 미선정 시 상시 운영 유지.', '행사일수 +α, 기준 시나리오의 가정 충족', '선정 대기 중. 3주차(10/12~18) 오늘끝딜이 포함된 주의 ROAS는 1,300%대', '선정 대기'),
 ('아넬라', '메타 스스', '차주까지', '할인 소구 소재 1종과 브랜딩 또는 제품 소구 소재 1종을 교체 투입.', '장기 유지 소재의 피로도 해소, 상시 구간 ROAS 개선', '4주차 메타 스스 ROAS 하락 (주간 리포트)', '소재 준비 중'),
 ('아넬라', '메타 카탈로그', '11월 1~2주', '카탈로그 상품 세트를 영업팀 주력 상품 기준으로 재구성하고 ROAS 200% 이하 구간 개선.', '카탈로그 효율 회복', '4주차 카탈로그 ROAS 189% (주간 리포트)', '제안'),
 ('아넬라', '메타 자사몰', '11월 1주', '개편된 시크릿링크 썸네일로 광고 소재 제작. 웰컴박스·기획전은 보류를 유지.', '자사몰 시크릿링크 ROAS 회복', '시크릿링크 ROAS가 주차 경과에 따라 하락 (주차별 추이)', '제작 예정'),
 ('아넬라', 'GFA 전환', '소재 검토 후', '신규 소구점 5종 테스트는 소재 검토를 마친 뒤 재개. 재개 전 소구점·타겟 가설을 정리.', '신규 소구점 발굴', '10/14~16 집행 시 구매 1건, 현재 HOLD', 'HOLD'),
 ('아넬라', '쇼핑검색', '넾다세일 기간(~11/8)', '넾다세일(10/26)부터 24시간 균등 노출로 원복 완료. 이후 시간대별 성과를 함께 확인. 과자(떡뻥) 캠페인은 롱테일 검색어 중심으로 정리.', '행사 기간 낮 시간대 유입·매출 기회 확보', '시간 제한 후 ROAS 490%→368%, 매출 -24% (시간제한_검증 시트). 시간대별 입찰비중 조정 기능은 확인되지 않음', '원복 완료'),
 ('아넬라', '파워링크', '상시', '현행 노출 시간 설정(노출 제외 0~3시, 17~24시)을 유지하고 매출 볼륨을 추가로 모니터링.', '효율 개선분(ROAS 224%→408%) 유지', '광고비 -58%에 매출 -22%. 10/21 저효율 키워드 감액이 같은 기간에 섞여 있음 (시간제한_검증 시트)', '진행 중'),
 ('로지오가닉', '메타 스스', '차주', '네이버 협력광고 신규 소재 기획안을 전달하고 소재를 교체. 10/24~25 이상치 원인을 먼저 확인.', '집계 신뢰도 확보, 메타 스스 효율 유지', '점검 시트 1번 (이틀 전환값이 스토어 결제의 수 배)', '확인 필요'),
 ('로지오가닉', 'GFA 전환', '11/1~11/8', '넾다세일 한정 20만원으로 웹전환 캠페인을 운영하고 종료 시 지속 여부 판단.', '행사 구간 추가 전환 확보', '10/26~28 3일 ROAS 111%, 표본 작음', '진행 중'),
 ('로지오가닉', '메타 자사몰', '예산 여유 시', '배도라지즙 시크릿링크 소재를 제작해 광고비 여유가 생기면 집행. 보리차 시크릿링크는 보류 유지.', '자사몰 매출 다변화', '광고비 여유 여부에 따라 집행. 시크릿링크 개편 썸네일 활용', '조건부'),
 ('로지오가닉', '네이버 검색', '11/9~', '파워링크 대표K·구매K 그룹 중심으로 저효율 키워드를 정리하고 넾다세일에 맞춰 일예산을 유동적으로 조정.', '파워링크·쇼핑검색 ROAS를 하한 근처로 회복', '10월 파워링크+쇼핑검색 합산 ROAS가 하한 미달 (성과분석 시트)', '진행 중'),
 ('공통', '12월', '11월 말', '12월 예산은 11월 수준을 유지. 연말 행사 일정이 영업팀에서 나오는 대로 목표 시트의 12월 행사일수를 확정.', '12월 목표 확정', '현재 12월 행사 일정 없음', '예정'),
]
for i, (b, m, when, what, why, basis, status) in enumerate(plan):
    r = 5 + i
    put(ep, f'A{r}', i + 1, al=CEN, border=True)
    put(ep, f'B{r}', b, al=CEN, border=True); put(ep, f'C{r}', m, al=CEN, border=True); put(ep, f'D{r}', when, al=CEN, border=True)
    put(ep, f'E{r}', what, al=LEFT, border=True); put(ep, f'F{r}', why, al=LEFT, border=True); put(ep, f'G{r}', basis, f_note, al=LEFT, border=True)
    put(ep, f'H{r}', status, f_bold, al=CEN, border=True)
    ep.row_dimensions[r].height = 56
ep.freeze_panes = 'A5'

# =====================================================================
# 요약
# =====================================================================
sm = sheet('요약', '1F3864')
widths(sm, {'A': 16, 'B': 15, 'C': 15, 'D': 13, 'E': 13, 'F': 13, 'G': 15, 'H': 15, 'I': 15, 'J': 13, 'K': 13})
put(sm, 'A1', '2026년 10월 광고 성과 및 11~12월 목표 · 아넬라 / 로지오가닉', f_title)
put(sm, 'A2', f'="집계 "&TEXT({S_GA}!B5,"yyyy.mm.dd")&" ~ "&TEXT({S_GA}!B6,"mm.dd")&" ("&{S_GA}!B7&"일) · 월말 예상은 10/26~28 일평균 기준 · 금액 단위 만원 · 네이버 VAT 미포함 · 매출은 매체 구매 귀속 기준"', f_note)

put(sm, 'A4', '1. 한눈에 보기', f_sec)
b1 = (f'="· 아넬라: 28일 ROAS "&TEXT({S_RS}!G31,"#,##0%")&"로 KPI 하한("&TEXT({S_RS}!C31,"0%")&")의 "&TEXT({S_RS}!H31,"0%")&". 월말 예상 ROAS "&TEXT({S_RS}!D38,"#,##0%")&"이나 예산은 일할 대비 "&TEXT({S_RS}!E{ra["tot"]},"0%")&"만 사용."')
b2 = (f'="· 로지오가닉: 28일 ROAS "&TEXT({S_RS}!G32,"#,##0%")&"로 하한("&TEXT({S_RS}!C32,"0%")&")의 "&TEXT({S_RS}!H32,"0%")&". 메타 스스 이상치 이틀을 빼도 "&TEXT({S_RS}!D46,"#,##0%")&"로 하한 상회."')
b3 = (f'="· 아넬라는 행사 구간 ROAS("&TEXT({S_WK}!I{wa["roas"]},"#,##0%")&")가 상시 구간("&TEXT({S_WK}!H{wa["roas"]},"#,##0%")&")의 "&TEXT({S_WK}!I{wa["roas"]}/{S_WK}!H{wa["roas"]},"0.0")&"배로, 11월 목표는 행사 일정에 좌우. 하한 달성 필요 행사일수는 약 "&TEXT({S_TG}!O{TGR["11월"]["judge"]["아넬라"]},"0.0")&"일(11월)."')
for i, f in enumerate([b1, b2, b3]):
    put(sm, f'A{5+i}', f, al=LEFT); sm.merge_cells(f'A{5+i}:K{5+i}'); sm.row_dimensions[5 + i].height = 30

put(sm, 'A9', '2. 10월 실적과 목표(소급) 대비', f_sec)
hdr(sm, 10, 1, ['브랜드', '광고비', '매출', 'ROAS', 'KPI 하한', 'ROAS 달성률', '월말 예상 ROAS', '월말 예상 매출 ÷ 목표', '구매(건)', '스토어 일반 순매출', '28일 판정'], height=44)
for i, b in enumerate(BRANDS):
    r = 11 + i; t = RSR[b]['tot']; w = WKR[b]
    put(sm, f'A{r}', b, f_bold, border=True)
    put(sm, f'B{r}', f"={S_RS}!E{31+i}", fmt=INT, border=True)
    put(sm, f'C{r}', f"={S_RS}!F{31+i}", fmt=INT, border=True)
    put(sm, f'D{r}', f"={S_RS}!G{31+i}", f_bold, PCT, border=True)
    put(sm, f'E{r}', f"={S_RS}!C{31+i}", fmt=PCT, border=True)
    put(sm, f'F{r}', f"={S_RS}!H{31+i}", f_bold, PCT, border=True)
    put(sm, f'G{r}', f"={S_RS}!D{38+i}", fmt=PCT, border=True)
    put(sm, f'H{r}', f"={S_RS}!F{38+i}", fmt=PCT, border=True)
    put(sm, f'I{r}', f"={S_RS}!J{t}", fmt=INT, border=True)
    put(sm, f'J{r}', f"={S_WK}!G{w['ss']}/10000", fmt=INT, border=True)
    put(sm, f'K{r}', f"={S_RS}!I{31+i}", f_bold, al=CEN, border=True)
put(sm, 'A13', '합계', f_bold, fill=fill_sub, border=True)
for c, src, fmt in [('B', 'B', INT), ('C', 'C', INT)]:
    put(sm, f'{c}13', f'={c}11+{c}12', f_bold, INT, fill_sub, border=True)
put(sm, 'D13', '=IF(B13=0,0,C13/B13)', f_bold, PCT, fill_sub, border=True)
put(sm, 'E13', f"={S_RS}!C33", f_bold, PCT, fill_sub, border=True)
put(sm, 'F13', f"={S_RS}!H33", f_bold, PCT, fill_sub, border=True)
put(sm, 'G13', f"={S_RS}!D40", f_bold, PCT, fill_sub, border=True)
put(sm, 'H13', f"={S_RS}!F40", f_bold, PCT, fill_sub, border=True)
put(sm, 'I13', '=I11+I12', f_bold, INT, fill_sub, border=True)
put(sm, 'J13', '=J11+J12', f_bold, INT, fill_sub, border=True)
put(sm, 'K13', f"={S_RS}!I33", f_bold, al=CEN, fill=fill_sub, border=True)
put(sm, 'A14', '10월 목표 부재로 4분기 KPI 보고안의 ROAS 하한(아넬라 850%, 로지 420%)과 10월 광고예산을 소급 적용. 스토어 일반 순매출은 공동구매 상품 제외 값.', f_note)
sm.conditional_formatting.add('K11:K13', FormulaRule(formula=['K11="달성"'], fill=PatternFill('solid', bgColor='C6E0B4', fgColor='C6E0B4')))
sm.conditional_formatting.add('K11:K13', FormulaRule(formula=['K11="근접"'], fill=PatternFill('solid', bgColor='FFE699', fgColor='FFE699')))
sm.conditional_formatting.add('K11:K13', FormulaRule(formula=['K11="미달"'], fill=PatternFill('solid', bgColor='F8CBAD', fgColor='F8CBAD')))

put(sm, 'A16', '3. 매체별 ROAS와 광고비 (10/1~10/28)', f_sec)
hdr(sm, 17, 1, ['매체', '아넬라 10월 예산', '아넬라 소진금액', '아넬라 소진율 (월 예산 대비)', '아넬라 매출', '아넬라 ROAS', '로지 10월 예산', '로지 소진금액', '로지 소진율 (월 예산 대비)', '로지 매출', '로지 ROAS'], height=44)
for k, m in enumerate(MEDIA7):
    r = 18 + k
    put(sm, f'A{r}', m if m != '브랜드검색' else '브랜드검색 (정액)', border=True)
    for (bname, cb, cc, cd, ce, cf) in [('아넬라', 'B', 'C', 'D', 'E', 'F'), ('로지오가닉', 'G', 'H', 'I', 'J', 'K')]:
        put(sm, f'{cb}{r}', f"={rsm(bname, m, 'B')}", fmt=INT, border=True)
        put(sm, f'{cc}{r}', f"={rsm(bname, m, 'D')}", fmt=INT, border=True)
        put(sm, f'{cd}{r}', f'=IF({cb}{r}=0,0,{cc}{r}/{cb}{r})', fmt=PCTD, border=True)
        put(sm, f'{ce}{r}', f"={rsm(bname, m, 'L')}", fmt=INT, border=True)
        put(sm, f'{cf}{r}', f"={rsm(bname, m, 'M')}", f_bold, PCTD, border=True)
put(sm, 'A25', '합계', f_bold, fill=fill_sub, border=True)
for (cb, cc, cd, ce, cf) in [('B', 'C', 'D', 'E', 'F'), ('G', 'H', 'I', 'J', 'K')]:
    for c in (cb, cc, ce):
        put(sm, f'{c}25', f'=SUM({c}18:{c}24)', f_bold, INT, fill_sub, border=True)
    put(sm, f'{cd}25', f'=IF({cb}25=0,0,{cc}25/{cb}25)', f_bold, PCT, fill_sub, border=True)
    put(sm, f'{cf}25', f'=IF({cc}25=0,0,{ce}25/{cc}25)', f_bold, PCT, fill_sub, border=True)
put(sm, 'A26', f'="· 금액 단위 만원. 소진율은 월 예산 대비이며 28/31 일할 기준으로는 아넬라 "&TEXT({S_RS}!E{ra["tot"]},"0%")&", 로지 "&TEXT({S_RS}!E{rr_["tot"]},"0%")&". 브랜드검색은 정액이라 예산을 정액(31일분)으로 표시하고 ROAS는 월 정액 대비 값. GFA 전환 로지오가닉은 10월 예산 0, 10/26부터 넾다세일 한정 집행."', f_note)

put(sm, 'A28', '4. 잘한 점 · 성과 이유 / 미달 · 보완', f_sec)
good = [
 f'="· 아넬라 파워링크 노출 시간 제한(10/14~): 행사 없는 두 주 비교 시 광고비 "&TEXT({S_TV}!E7,"+0%;-0%")&"에도 구매 "&{S_TV}!B10&"건에서 "&{S_TV}!C10&"건으로 유지, ROAS "&TEXT({S_TV}!B13,"#,##0%")&"에서 "&TEXT({S_TV}!C13,"#,##0%")&"로 상승."',
 f'="· 아넬라 저효율 키워드·지면 정리와 시간 제한으로 일할 예산의 "&TEXT({S_RS}!E{ra["tot"]},"0%")&"만 집행하고도 ROAS가 하한에 근접("&TEXT({S_RS}!H31,"0%")&"). 상시 구간에서는 쇼핑검색(ROAS "&TEXT({wkm("아넬라","쇼핑검색","H")},"#,##0%")&")이 파워링크("&TEXT({wkm("아넬라","파워링크","H")},"#,##0%")&")보다 효율 높음."',
 f'="· 메타 시크릿링크가 자사몰 매출의 주축. 시크릿특가 상품이 자사몰 매출의 아넬라 "&TEXT({S_JM}!H7,"0%")&", 로지 "&TEXT({S_JM}!H8,"0%")&"를 차지, 시크릿특가 매출 기준 ROAS는 아넬라 "&TEXT({S_JM}!G15,"#,##0%")&", 로지 "&TEXT({S_JM}!G16,"#,##0%")&"."',
 f'="· 로지오가닉은 행사 영향이 작은 가운데 하한 상회(ROAS "&TEXT({S_RS}!G32,"#,##0%")&", 이상치 제외 "&TEXT({S_RS}!D46,"#,##0%")&"), 메타 자사몰 과자 시크릿링크는 10월 내내 안정적("&TEXT({rsm("로지오가닉","메타 자사몰","M")},"#,##0%")&")."',
]
bad = [
 f'="· 아넬라 상시 구간 ROAS "&TEXT({S_WK}!H{wa["roas"]},"#,##0%")&"로 하한("&TEXT({S_RS}!C31,"0%")&")의 절반 수준. 4주차(10/19~25) 전 매체 ROAS 하락. 보완: 상시 소재 교체, 쇼핑검색 24시간 원복 후 성과 확인, 저효율 정리 효과 검증."',
 f'="· 아넬라 메타 자사몰·카탈로그 효율 하락. 웰컴박스·기획전은 보류 유지, 시크릿링크는 개편 썸네일로 교체, 카탈로그는 주력 상품 중심으로 재구성."',
 f'="· 로지오가닉 파워링크+쇼핑검색 합산 ROAS "&TEXT({S_AN}!E{ANR["파워링크·쇼핑검색 효율"]},"#,##0%")&"로 하한 미달. 보완: 파워링크 대표K·구매K 점검, 11/8 이후 상시 구간으로 감액 효과 검증."',
 f'="· 로지오가닉 메타 스스 10/24~25 전환값이 같은 이틀 스토어 결제의 "&TEXT({S_CK}!D15,"0.0")&"배. 집계 이슈 가능성으로 확인 전까지 목표 산정에서 제외."',
]
for i, f in enumerate(good):
    put(sm, f'A{29+i}', f, al=LEFT); sm.merge_cells(f'A{29+i}:K{29+i}'); sm.row_dimensions[29 + i].height = 30
    sm[f'A{29+i}'].fill = PatternFill('solid', fgColor='E2EFDA')
for i, f in enumerate(bad):
    put(sm, f'A{33+i}', f, al=LEFT); sm.merge_cells(f'A{33+i}:K{33+i}'); sm.row_dimensions[33 + i].height = 30
    sm[f'A{33+i}'].fill = PatternFill('solid', fgColor='FCE4D6')

put(sm, 'A38', '5. 11~12월 목표 (광고비는 확정 예산, 목표 매출은 행사일수 가정에 따른 시나리오)', f_sec)
hdr(sm, 39, 1, ['브랜드', '월', '광고비', '목표 ROAS 보수', '목표 ROAS 기준', '목표 ROAS 도전', '목표 매출 보수', '목표 매출 기준', '목표 매출 도전', 'KPI 하한', '하한 달성 필요 행사일수'], height=44)
r = 40
for b in BRANDS:
    for month in ['11월', '12월']:
        s, e, tot = TGR[month]['rows'][b]
        jr = TGR[month]['judge'][b]
        put(sm, f'A{r}', b, f_bold, border=True); put(sm, f'B{r}', month, al=CEN, border=True)
        put(sm, f'C{r}', f"={S_TG}!K{tot}", fmt=INT, border=True)
        for j in range(3):
            put(sm, f'{"DEF"[j]}{r}', f"={S_TG}!{'LMN'[j]}{tot}", f_bold, PCT, border=True)
            put(sm, f'{"GHI"[j]}{r}', f"={S_TG}!{'OPQ'[j]}{tot}", fmt=INT, border=True)
        put(sm, f'J{r}', f"={S_TG}!B{jr}", fmt=PCT, border=True)
        put(sm, f'K{r}', f"={S_TG}!O{jr}", fmt='0.0"일"', border=True)
        r += 1
put(sm, f'A{r}', f'="가정: 11월 행사일수 "&{S_TG}!B6&"/"&{S_TG}!C6&"/"&{S_TG}!D6&"일(보수/기준/도전), 12월 "&{S_TG}!B7&"/"&{S_TG}!C7&"/"&{S_TG}!D7&"일. 11월 확정 일정은 넾다세일 8일뿐, 나머지는 영업팀 일정에 따라 수정."', f_note)
put(sm, f'A{r+1}', f'="스마트스토어 일반상품 순매출 목표(기준, 공구 제외): 아넬라 11월 "&TEXT({S_TG}!J{ssr[("아넬라","11월")]},"#,##0")&"만원, 12월 "&TEXT({S_TG}!J{ssr[("아넬라","12월")]},"#,##0")&"만원 / 로지오가닉 11월 "&TEXT({S_TG}!J{ssr[("로지오가닉","11월")]},"#,##0")&"만원, 12월 "&TEXT({S_TG}!J{ssr[("로지오가닉","12월")]},"#,##0")&"만원"', f_note)
put(sm, f'A{r+2}', f'="목표 매출은 예산 100% 사용 가정(10월 소진율: 아넬라 "&TEXT({S_RS}!E{ra["tot"]},"0%")&", 로지 "&TEXT({S_RS}!E{rr_["tot"]},"0%")&"). 아넬라는 기준 시나리오에서도 하한 미달로, 하한 유지 시 행사일 확보 필요. 행사 부족 시 하한 조정 논의 필요."', f_note)
SM_R = r + 2

put(sm, f'A{SM_R+3}', '6. 확인이 필요한 사항', f_sec)
confirm = ['· 블랙프라이데이·12월 행사 일정 (영업팀). 확정 시 목표 시트의 행사일수만 수정하면 목표 자동 재계산.',
           '· 로지오가닉 메타 스스 10/24~25 전환 집계 (점검 시트 1번).',
           '· 10/29~10/31 확정 실적. 확정 시 월말 예상을 실적으로 교체.',
           '· 네이버 광고비의 VAT 기준이 예산표(VAT 미포함)와 같은지 확인.',
           '· 아넬라 GFA 전환의 귀속 지연(노출 0인 날의 전환). 확인 전까지 해당 매체 행사 구간 ROAS는 높게 나올 가능성.']
for i, t in enumerate(confirm):
    put(sm, f'A{SM_R+4+i}', t, al=LEFT); sm.merge_cells(f'A{SM_R+4+i}:K{SM_R+4+i}')


b7 = SM_R + 4 + len(confirm) + 1
put(sm, f'A{b7}', '7. 자사몰 매출 (상품명에 "시크릿 특가" = 메타 시크릿링크 유입)', f_sec)
hdr(sm, b7 + 1, 1, ['브랜드', '자사몰 매출', '시크릿특가 매출', '시크릿특가 비중', '시크릿링크 광고비', '메타 보고 매출', '보고 ROAS', '시크릿특가 실매출 ROAS'], height=44)
for i, b in enumerate(BRANDS):
    rr = b7 + 2 + i; jr = 7 + i; mr = 15 + i
    put(sm, f'A{rr}', b, f_bold, border=True)
    put(sm, f'B{rr}', f"={S_JM}!F{jr}", fmt=INT, border=True)
    put(sm, f'C{rr}', f"={S_JM}!G{jr}", fmt=INT, border=True)
    put(sm, f'D{rr}', f"={S_JM}!H{jr}", fmt=PCT, border=True)
    put(sm, f'E{rr}', f"={S_JM}!B{mr}", fmt=INT, border=True)
    put(sm, f'F{rr}', f"={S_JM}!C{mr}", fmt=INT, border=True)
    put(sm, f'G{rr}', f"={S_JM}!F{mr}", fmt=PCT, border=True)
    put(sm, f'H{rr}', f"={S_JM}!G{mr}", f_bold, PCT, border=True)
put(sm, f'A{b7+4}', '자사몰 매출은 파일의 판매합계 기준(만원, 실결제 기준), 기간은 10/1~10/28. 아넬라·로지 전품목 맛보기팩(공통)은 브랜드별 합계에서 제외.', f_note)
sm.merge_cells(f'A{b7+4}:K{b7+4}'); sm.row_dimensions[b7 + 4].height = 28

sm.freeze_panes = 'A4'



# =====================================================================
# 시간제한_검증
# =====================================================================
tv = sheet('시간제한_검증', '00B0F0')
widths(tv, {'A': 26, 'B': 16, 'C': 16, 'D': 14, 'E': 12, 'F': 4, 'G': 60})
put(tv, 'A1', '아넬라 네이버 검색광고 노출 시간 제한(10/14) 효과 검증', f_title)
put(tv, 'A2', '파워링크·쇼핑검색 시간대 지정 노출 결과(노출 제외 0~3시, 17~24시, 두 매체 동일, 10/14 시작). 쇼핑검색은 10/26 넾다세일부터 24시간 균등 노출로 원복. 수치는 데이터_광고 시트 SUMIFS 수식이며 대행사 보고 수치와 일치.', f_note)
tv.merge_cells('A2:G2')
put(tv, 'A3', f'="비교기간 "&TEXT({S_GA}!B51,"m/d")&"~"&TEXT({S_GA}!C51,"m/d")&" (10월 2주차, 제한 전) vs 기준기간 "&TEXT({S_GA}!B53,"m/d")&"~"&TEXT({S_GA}!C53,"m/d")&" (10월 4주차, 제한 후). 두 기간 모두 상시 구간."', f_note)
tv.merge_cells('A3:G3')
metrics = [('광고비(원)', 'E', WON), ('클릭', 'G', WON), ('구매(건)', 'H', WON), ('매출(원)', 'I', WON)]
TVR = {}
def tv_block(r0, media):
    put(tv, f'A{r0}', f'아넬라 {media}', f_sec)
    hdr(tv, r0 + 1, 1, ['지표', '비교기간 (제한 전)', '기준기간 (제한 후)', '증감', '증감률'], height=30)
    rows = ['광고비(원)', '클릭', 'CPC(원)', '구매(건)', '전환율 (구매÷클릭)', '매출(원)', 'ROAS']
    R = {n: r0 + 2 + i for i, n in enumerate(['cost', 'clk', 'cpc', 'buy', 'cvr', 'rev', 'roas'])}
    TVR[media] = R
    for i, t in enumerate(rows):
        put(tv, f'A{r0+2+i}', t, border=True)
    for col, prow in [('B', 51), ('C', 53)]:
        crit = f'{adr("A")},"아넬라",{adr("B")},"{media}",{adr("D")},{S_GA}!$A${prow}'
        for key, dcol in [('cost', 'E'), ('clk', 'G'), ('buy', 'H'), ('rev', 'I')]:
            put(tv, f'{col}{R[key]}', f'=SUMIFS({adr(dcol)},{crit})', fmt=WON, border=True)
        put(tv, f'{col}{R["cpc"]}', f'=IF({col}{R["clk"]}=0,0,{col}{R["cost"]}/{col}{R["clk"]})', fmt=WON, border=True)
        put(tv, f'{col}{R["cvr"]}', f'=IF({col}{R["clk"]}=0,0,{col}{R["buy"]}/{col}{R["clk"]})', fmt='0.00%', border=True)
        put(tv, f'{col}{R["roas"]}', f'=IF({col}{R["cost"]}=0,0,{col}{R["rev"]}/{col}{R["cost"]})', f_bold, PCT, border=True)
    for key in ['cost', 'clk', 'cpc', 'buy', 'rev']:
        r = R[key]
        put(tv, f'D{r}', f'=C{r}-B{r}', fmt=WON, border=True)
        put(tv, f'E{r}', f'=IF(B{r}=0,0,C{r}/B{r}-1)', fmt='+0%;-0%;0%', border=True)
    put(tv, f'D{R["cvr"]}', f'=(C{R["cvr"]}-B{R["cvr"]})*100', fmt='+0.00"%p";-0.00"%p"', border=True)
    put(tv, f'E{R["cvr"]}', '', border=True)
    put(tv, f'D{R["roas"]}', f'=(C{R["roas"]}-B{R["roas"]})*100', fmt='+0"%p";-0"%p"', border=True)
    put(tv, f'E{R["roas"]}', '', border=True)
    return r0 + 10
n1 = tv_block(5, '파워링크')
n2 = tv_block(n1 + 1, '쇼핑검색')
r = n2 + 1
put(tv, f'A{r}', '참고: 24시간 원복 후 3일 (10/26~10/28, 넾다세일 겹침) 아넬라 쇼핑검색 ROAS', f_bold, border=True)
put(tv, f'C{r}', f'=IF(SUMIFS({adr("E")},{adr("A")},"아넬라",{adr("B")},"쇼핑검색",{adr("D")},{S_GA}!$A$54)=0,0,SUMIFS({adr("I")},{adr("A")},"아넬라",{adr("B")},"쇼핑검색",{adr("D")},{S_GA}!$A$54)/SUMIFS({adr("E")},{adr("A")},"아넬라",{adr("B")},"쇼핑검색",{adr("D")},{S_GA}!$A$54))', f_bold, PCT, border=True)
tv.merge_cells(f'A{r}:B{r}')
put(tv, f'G{r}', '원복과 넾다세일이 같은 날 시작되어 원복 효과 분리 불가. 11/8 이후 상시 구간과 비교 필요.', f_note)
r += 2
put(tv, f'A{r}', '판단과 제안', f_sec)
hdr(tv, r + 1, 1, ['매체', '판단', '', '', ''], height=24)
tv.merge_cells(f'B{r+1}:E{r+1}')
tv[f'G{r+1}'].value = '근거 · 확인할 점'
put(tv, f'G{r+1}', '근거 · 확인할 점', f_hdr, fill=fill_hdr, al=CEN, border=True)
verdict = [
 ('파워링크', '현행 시간 설정 유지. 매출 볼륨을 추가로 모니터링.',
  '광고비·클릭 약 58% 감소, 전환율 상승, ROAS 224%→408%. 매출은 -22%로 유입 감소 폭보다 작아 볼륨은 줄었지만 효율은 개선. '
  '확인 사항: 기준기간에 10/21 저효율 키워드 일예산 감액이 함께 반영되어 시간 제한만의 효과로 보기 어려움.'),
 ('쇼핑검색', '넾다세일부터(10/26) 24시간 균등 노출로 원복 완료. 이후 시간대별 성과를 함께 확인.',
  '유입 규모(광고비·클릭 약 +1%)는 유지됐지만 전환율, 매출(-24%), ROAS(490%→368%)가 하락. 노출 제외 시간에도 구매가 발생. '
  '넾다세일로 검색·구매 수요가 커질 가능성이 있어 낮 시간대까지 확보. 시간대별 입찰 비중 조정 기능은 확인되지 않음. '
  '확인 사항: 시간대별 구매 수치는 대행사 보고 기준이며 이 워크북에는 시간대 데이터 없음. 10/26 이후는 원복과 넾다세일이 겹쳐 원복 효과 분리 불가.'),
]
for i, (m, v, why) in enumerate(verdict):
    rr = r + 2 + i
    put(tv, f'A{rr}', m, f_bold, al=LEFT, border=True)
    put(tv, f'B{rr}', v, al=LEFT, border=True); tv.merge_cells(f'B{rr}:E{rr}')
    put(tv, f'G{rr}', why, f_note, al=LEFT, border=True)
    tv.row_dimensions[rr].height = 96
put(tv, f'A{r+5}', '출처: 노출 시간 제한 결과 정리는 대행사 코멘트(사용자 전달, 2026-10-29). 시간 설정은 10/14 시작. 노출 제외 시간대는 0~3시, 17~24시로 파워링크·쇼핑검색 동일(사용자 확인).', f_note)
tv.merge_cells(f'A{r+5}:G{r+5}')
tv.freeze_panes = 'A5'


# =====================================================================
# 예산
# =====================================================================
bg = sheet('예산', '70AD47')
widths(bg, {'A': 22, 'B': 12, 'C': 12, 'D': 11, 'E': 10, 'F': 13, 'G': 14, 'H': 13, 'I': 64})
put(bg, 'A1', '10월·11월 광고 예산 (만원, 네이버 VAT 미포함)', f_title)
put(bg, 'A2', '예산은 가정 시트 입력값 참조. 12월은 11월 유지 가정. 10월 집행은 10/1~10/28(28일) 광고비이며, 일할 대비 소진율은 10월 예산을 28/31로 환산한 값 대비.', f_note)
bg.merge_cells('A2:I2'); bg.row_dimensions[2].height = 28
reasons = {
 ('아넬라', '파워링크'): '변동 없음. 10/14 노출 시간 제한 유지(노출 제외 0~3시, 17~24시), 10/21 저효율 키워드 일예산 감액.',
 ('아넬라', '쇼핑검색'): '변동 없음. 10/26 넾다세일부터 24시간 균등 노출로 원복.',
 ('아넬라', 'GFA 카탈로그'): '변동 없음.',
 ('아넬라', 'GFA 전환'): '250만원에서 150만원으로 감액. 신규 소구점 5종 테스트가 소재 검토로 HOLD 중.',
 ('아넬라', '메타 스스'): '변동 없음. 넾다세일부터 인스타그램 지면 고정, 광고 상품은 영업팀 주력 상품으로 전환.',
 ('아넬라', '메타 자사몰'): '250만원에서 180만원으로 축소. 웰컴박스·기획전 보류.',
 ('로지오가닉', '파워링크'): '변동 없음. 10/21 저효율 키워드 일예산 감액.',
 ('로지오가닉', '쇼핑검색'): '변동 없음. 10/26 무전환 지면 노출 제외.',
 ('로지오가닉', 'GFA 카탈로그'): '변동 없음.',
 ('로지오가닉', 'GFA 전환'): '11월 신설 20만원, 넾다세일 한정. 12월은 11월 유지 가정이라 같은 값이 들어가 있어 사용처 확정 필요.',
 ('로지오가닉', '메타 스스'): '250만원에서 215만원으로 약간 감액. 넾다세일 영향.',
 ('로지오가닉', '메타 자사몰'): '180만원에서 150만원으로 감액. 보리차 시크릿링크 보류.',
}
BGR = {}
def bg_block(r0, brand):
    put(bg, f'A{r0}', brand, f_sec)
    hdr(bg, r0 + 1, 1, ['매체', '10월 예산', '11월 예산', '증감', '증감률', '12월 (11월 유지)', '10월 집행 (28일)', '일할 대비 소진율', '변경 사유 · 운영 메모'], height=32)
    first = r0 + 2
    for k, m in enumerate(MEDIA6):
        r = first + k
        crit = f"{S_GA}!$A$16:$A$27,$A${r0},{S_GA}!$B$16:$B$27,$A{r}"
        put(bg, f'A{r}', m, border=True)
        put(bg, f'B{r}', f"=SUMIFS({S_GA}!$C$16:$C$27,{crit})", fmt=INT, border=True)
        put(bg, f'C{r}', f"=SUMIFS({S_GA}!$D$16:$D$27,{crit})", fmt=INT, border=True)
        put(bg, f'F{r}', f"=SUMIFS({S_GA}!$E$16:$E$27,{crit})", fmt=INT, border=True)
        put(bg, f'G{r}', f"={S_RS}!D{RSR[brand]['first'] + MEDIA7.index(m)}", fmt=MAN, border=True)
        put(bg, f'I{r}', reasons[(brand, m)], al=LEFT, border=True)
        bg.row_dimensions[r].height = 30
    tot = first + len(MEDIA6)
    last = tot - 1
    put(bg, f'A{tot}', f'{brand} 합계', f_bold, fill=fill_sub, border=True)
    for c in 'BCFG':
        put(bg, f'{c}{tot}', f'=SUM({c}{first}:{c}{last})', f_bold, INT if c != 'G' else MAN, fill_sub, border=True)
    for r in range(first, tot + 1):
        bold = f_bold if r == tot else f_norm
        fill = fill_sub if r == tot else None
        put(bg, f'D{r}', f'=C{r}-B{r}', bold, '+#,##0;-#,##0;0', fill, border=True)
        put(bg, f'E{r}', f'=IF(B{r}=0,0,C{r}/B{r}-1)', bold, '+0%;-0%;0%', fill, border=True)
        put(bg, f'H{r}', f"=IF(B{r}=0,0,G{r}/(B{r}*{S_GA}!$B$7/{S_GA}!$B$8))", bold, PCTD, fill, border=True)
    put(bg, f'I{tot}', '', fill=fill_sub, border=True)
    # brand search reference row
    bs = tot + 1
    put(bg, f'A{bs}', '(참고) 브랜드검색 정액', f_note, border=True)
    put(bg, f'B{bs}', f"={S_GA}!$B$38*{S_GA}!$B$8/10000", f_note, MAN, border=True)
    put(bg, f'C{bs}', f"={S_GA}!$B$38*{S_GA}!$B$10/10000", f_note, MAN, border=True)
    put(bg, f'D{bs}', f'=C{bs}-B{bs}', f_note, '+#,##0.0;-#,##0.0;0.0', border=True)
    put(bg, f'E{bs}', '', border=True)
    put(bg, f'F{bs}', f"={S_GA}!$B$38*{S_GA}!$B$11/10000", f_note, MAN, border=True)
    put(bg, f'G{bs}', f"={S_RS}!D{RSR[brand]['first'] + MEDIA7.index('브랜드검색')}", f_note, MAN, border=True)
    put(bg, f'H{bs}', '', border=True)
    put(bg, f'I{bs}', '예산표에 미포함된 정액 상품(월 88만원을 일수로 안분).', f_note, al=LEFT, border=True)
    BGR[brand] = dict(first=first, tot=tot, bs=bs)
    return bs + 2
n = bg_block(4, '아넬라')
n = bg_block(n, '로지오가닉')
r = n
put(bg, f'A{r}', '전체 합계', f_bold, fill=fill_sub, border=True)
for c in 'BCFG':
    put(bg, f'{c}{r}', f"={c}{BGR['아넬라']['tot']}+{c}{BGR['로지오가닉']['tot']}", f_bold, INT if c != 'G' else MAN, fill_sub, border=True)
put(bg, f'D{r}', f'=C{r}-B{r}', f_bold, '+#,##0;-#,##0;0', fill_sub, border=True)
put(bg, f'E{r}', f'=IF(B{r}=0,0,C{r}/B{r}-1)', f_bold, '+0%;-0%;0%', fill_sub, border=True)
put(bg, f'H{r}', f"=IF(B{r}=0,0,G{r}/(B{r}*{S_GA}!$B$7/{S_GA}!$B$8))", f_bold, PCTD, fill_sub, border=True)
put(bg, f'I{r}', '', fill=fill_sub, border=True)
BG_TOT = r
put(bg, f'A{r+1}', '검증: 제시 합계와 차이(0이면 일치)', f_note, border=True)
put(bg, f'B{r+1}', f"={S_GA}!C34", f_note, INT, border=True)
put(bg, f'C{r+1}', f"={S_GA}!D34", f_note, INT, border=True)

r += 3
put(bg, f'A{r}', '4분기 KPI 보고안 광고예산과 비교 (브랜드검색 정액 포함 기준)', f_sec)
hdr(bg, r + 1, 1, ['브랜드', 'KPI안 10월', '실행 10월', '차이', '', 'KPI안 11월', '실행 11월', '차이', 'KPI안 12월 / 실행 12월 / 차이'], height=32)
for i, b in enumerate(BRANDS):
    rr = r + 2 + i; t = BGR[b]
    col = 'B' if b == '아넬라' else 'C'
    put(bg, f'A{rr}', b, f_bold, border=True)
    put(bg, f'B{rr}', f"={S_GA}!{col}44", fmt=INT, border=True)
    put(bg, f'C{rr}', f"=B{t['tot']}+B{t['bs']}", fmt=MAN, border=True)
    put(bg, f'D{rr}', f'=C{rr}-B{rr}', fmt='+#,##0.0;-#,##0.0;0.0', border=True)
    put(bg, f'E{rr}', '', border=True)
    put(bg, f'F{rr}', f"={S_GA}!{col}45", fmt=INT, border=True)
    put(bg, f'G{rr}', f"=C{t['tot']}+C{t['bs']}", fmt=MAN, border=True)
    put(bg, f'H{rr}', f'=G{rr}-F{rr}', fmt='+#,##0.0;-#,##0.0;0.0', border=True)
    put(bg, f'I{rr}', f'=TEXT({S_GA}!{col}46,"#,##0")&" / "&TEXT(F{t["tot"]}+F{t["bs"]},"#,##0.0")&" / "&TEXT(F{t["tot"]}+F{t["bs"]}-{S_GA}!{col}46,"+#,##0.0;-#,##0.0;0.0")', al=LEFT, border=True)
put(bg, f'A{r+4}', 'KPI안의 10월 예산(1,300 / 900)은 실행 예산에 브랜드검색 정액을 더한 값과 거의 일치. 아넬라 11~12월은 KPI안 대비 낮고, 로지오가닉은 KPI안 대비 높게 편성.', f_note)
bg.merge_cells(f'A{r+4}:I{r+4}')
put(bg, f'A{r+6}', '운영 메모 (사용자 전달 사항)', f_sec)
memos = ['넾다세일(10/26~11/8, 아넬라·로지 모두)부터 메타 인스타그램 지면 고정 운영.',
         '10월 중순부터 효율이 낮은 키워드·지면 OFF 또는 감액 중.',
         '넾다세일부터 메타 광고 상품은 상시 상품(아넬라 퓨레 6+6, 로지 쌀과자 8봉)이 아니라 영업팀 주력 상품으로 운영.',
         '12월 예산은 11월 예산 유지 예정. 블랙프라이데이 등 행사 일정 확정 시 조정.']
for i, t in enumerate(memos):
    put(bg, f'A{r+7+i}', f'· {t}', al=LEFT); bg.merge_cells(f'A{r+7+i}:I{r+7+i}')
bg.freeze_panes = 'A4'


# =====================================================================
# 예산대비_집행
# =====================================================================
bu = sheet('예산대비_집행', '70AD47')
widths(bu, {'A': 24, 'B': 11, 'C': 11, 'D': 12, 'E': 11, 'F': 11, 'G': 11, 'H': 12, 'I': 11, 'J': 9, 'K': 11, 'L': 9, 'M': 12, 'N': 11, 'O': 11})
put(bu, 'A1', '10월 예산 대비 집행 광고비·매출 (매체별)', f_title)
put(bu, 'A2', '금액 단위 만원(네이버 VAT 미포함). 집행·매출은 10/1~10/28(28일) 실적, 월말 예상은 28일 누계에 10/26~28 일평균 × 잔여일수를 더한 값. 매출은 매체 보고 구매 귀속 매출.', f_note)
bu.merge_cells('A2:O2'); bu.row_dimensions[2].height = 28
put(bu, 'A3', '소진 판정 기준 (입력)', f_bold)
put(bu, 'B3', 1.05, f_in, PCT, fill_in, RIGHT, True); put(bu, 'C3', '초과 이상', f_note)
put(bu, 'D3', 0.8, f_in, PCT, fill_in, RIGHT, True); put(bu, 'E3', '미소진 미만 (월말 예상 소진율 기준)', f_note)
BUR = {}
def bu_block(r0, brand):
    put(bu, f'A{r0}', brand, f_sec)
    hdr(bu, r0 + 1, 1, ['매체', '월 예산', '일할 예산 (28일)', '집행 광고비', '일할 대비 소진율', '월 예산 대비 소진율', '잔여 예산', '월말 예상 광고비', '월말 예상 소진율', '구매 (건)', '매출', 'ROAS', '월말 예상 매출', '예상 매출 ÷ 월 예산', '소진 판정'], height=44)
    first = r0 + 2
    meds = MEDIA6 + ['브랜드검색']
    for k, m in enumerate(meds):
        r = first + k
        crit = f'{adr("A")},$A${r0},{adr("B")},$A{r}'
        crit5 = crit + f',{adr("D")},{S_GA}!$A$54'
        put(bu, f'A{r}', m if m != '브랜드검색' else '브랜드검색 (정액, 예산표 외)', border=True)
        if m == '브랜드검색':
            put(bu, f'B{r}', f"={S_GA}!$B$38*{S_GA}!$B$8/10000", fmt=MAN, border=True)
            crit = f'{adr("A")},$A${r0},{adr("B")},"브랜드검색"'
            crit5 = crit + f',{adr("D")},{S_GA}!$A$54'
        else:
            put(bu, f'B{r}', f"=SUMIFS({S_GA}!$C$16:$C$27,{S_GA}!$A$16:$A$27,$A${r0},{S_GA}!$B$16:$B$27,$A{r})", fmt=MAN, border=True)
        put(bu, f'D{r}', f'=SUMIFS({adr("E")},{crit})/10000', fmt=MAN, border=True)
        put(bu, f'J{r}', f'=SUMIFS({adr("H")},{crit})', fmt=INT, border=True)
        put(bu, f'K{r}', f'=SUMIFS({adr("I")},{crit})/10000', fmt=MAN, border=True)
        put(bu, f'H{r}', f"=D{r}+SUMIFS({adr('E')},{crit5})/10000/{S_GA}!$D$54*{S_GA}!$B$9", fmt=MAN, border=True)
        put(bu, f'M{r}', f"=K{r}+SUMIFS({adr('I')},{crit5})/10000/{S_GA}!$D$54*{S_GA}!$B$9", fmt=MAN, border=True)
    t1 = first + len(meds)
    t2 = t1 + 1
    put(bu, f'A{t1}', f'{brand} 합계 (예산표 기준, 브랜드검색 제외)', f_bold, fill=fill_sub, border=True)
    put(bu, f'A{t2}', f'{brand} 합계 (브랜드검색 포함)', f_bold, fill=fill_sub, border=True)
    for c in 'BDHJKM':
        fmt = INT if c == 'J' else MAN
        put(bu, f'{c}{t1}', f'=SUM({c}{first}:{c}{first+len(MEDIA6)-1})', f_bold, fmt, fill_sub, border=True)
        put(bu, f'{c}{t2}', f'={c}{t1}+{c}{t1-1}', f_bold, fmt, fill_sub, border=True)
    for r in range(first, t2 + 1):
        tot = r >= t1
        bold = f_bold if tot else f_norm
        fill = fill_sub if tot else None
        put(bu, f'C{r}', f"=B{r}*{S_GA}!$B$7/{S_GA}!$B$8", bold, MAN, fill, border=True)
        put(bu, f'E{r}', f'=IF(C{r}=0,0,D{r}/C{r})', bold, PCTD, fill, border=True)
        put(bu, f'F{r}', f'=IF(B{r}=0,0,D{r}/B{r})', bold, PCTD, fill, border=True)
        put(bu, f'G{r}', f'=B{r}-D{r}', bold, '#,##0.0;-#,##0.0;0.0', fill, border=True)
        put(bu, f'I{r}', f'=IF(B{r}=0,0,H{r}/B{r})', bold, PCTD, fill, border=True)
        put(bu, f'L{r}', f'=IF(D{r}=0,0,K{r}/D{r})', f_bold, PCT, fill, border=True)
        put(bu, f'N{r}', f'=IF(B{r}=0,0,M{r}/B{r})', bold, '0.0"배"', fill, border=True)
        put(bu, f'O{r}', f'=IF(B{r}=0,"예산 없음",IF(I{r}>=$B$3,"초과",IF(I{r}<$D$3,"미소진","적정")))', f_bold, al=CEN, fill=fill, border=True)
    BUR[brand] = dict(first=first, t1=t1, t2=t2)
    return t2 + 2
n = bu_block(5, '아넬라')
n = bu_block(n, '로지오가닉')
r = n
put(bu, f'A{r}', '전체 합계 (브랜드검색 포함)', f_bold, fill=fill_sub, border=True)
for c in 'BDHJKM':
    fmt = INT if c == 'J' else MAN
    put(bu, f'{c}{r}', f"={c}{BUR['아넬라']['t2']}+{c}{BUR['로지오가닉']['t2']}", f_bold, fmt, fill_sub, border=True)
put(bu, f'C{r}', f"=B{r}*{S_GA}!$B$7/{S_GA}!$B$8", f_bold, MAN, fill_sub, border=True)
put(bu, f'E{r}', f'=IF(C{r}=0,0,D{r}/C{r})', f_bold, PCT, fill_sub, border=True)
put(bu, f'F{r}', f'=IF(B{r}=0,0,D{r}/B{r})', f_bold, PCT, fill_sub, border=True)
put(bu, f'G{r}', f'=B{r}-D{r}', f_bold, '#,##0.0;-#,##0.0;0.0', fill_sub, border=True)
put(bu, f'I{r}', f'=IF(B{r}=0,0,H{r}/B{r})', f_bold, PCT, fill_sub, border=True)
put(bu, f'L{r}', f'=IF(D{r}=0,0,K{r}/D{r})', f_bold, PCT, fill_sub, border=True)
put(bu, f'N{r}', f'=IF(B{r}=0,0,M{r}/B{r})', f_bold, '0.0"배"', fill_sub, border=True)
put(bu, f'O{r}', f'=IF(I{r}>=$B$3,"초과",IF(I{r}<$D$3,"미소진","적정"))', f_bold, al=CEN, fill=fill_sub, border=True)
BU_TOT = r
for i, b in enumerate(BRANDS):
    t = BUR[b]['t1']
    put(bu, f'A{r+2+i}', f'="· {b}: 월 예산 "&TEXT(B{t},"#,##0")&"만원 중 28일간 "&TEXT(D{t},"#,##0")&"만원 집행(일할 대비 "&TEXT(E{t},"0%")&"), 월말 예상 소진율 "&TEXT(I{t},"0%")&". 집행 대비 매출 "&TEXT(K{t},"#,##0")&"만원(ROAS "&TEXT(L{t},"#,##0%")&", 브랜드검색 제외)."', al=LEFT)
    bu.merge_cells(f'A{r+2+i}:O{r+2+i}')
put(bu, f'A{r+4}', '브랜드검색은 정액 상품이라 월 예산(88만원×31/30)과 집행이 거의 같고 매출이 커서 브랜드검색 포함 합계 ROAS가 높게 산출. 매체별 예산 소진 판단은 예산표 기준(브랜드검색 제외) 합계 기준.', f_note)
bu.merge_cells(f'A{r+4}:O{r+4}')
put(bu, f'A{r+5}', '월말 예상은 넾다세일 구간(10/26~28) 일평균을 3일 연장한 값. GFA 전환(로지)은 10월 예산이 0이라 소진율 산출 불가, 10/26부터 넾다세일 한정 집행.', f_note)
bu.merge_cells(f'A{r+5}:O{r+5}')
bu.freeze_panes = 'B5'
bu.conditional_formatting.add('O5:O60', FormulaRule(formula=['O5="초과"'], fill=PatternFill('solid', bgColor='F8CBAD', fgColor='F8CBAD')))
bu.conditional_formatting.add('O5:O60', FormulaRule(formula=['O5="미소진"'], fill=PatternFill('solid', bgColor='FFE699', fgColor='FFE699')))
bu.conditional_formatting.add('O5:O60', FormulaRule(formula=['O5="적정"'], fill=PatternFill('solid', bgColor='C6E0B4', fgColor='C6E0B4')))


# =====================================================================
# 데이터_자사몰 / 자사몰_매출
# =====================================================================
jm_raw = pd.read_csv(os.path.join(HERE, '..', 'source', 'jasamol_sales_2026-10.csv'), encoding='utf-8-sig')
dj = sheet('데이터_자사몰', 'A6A6A6')
widths(dj, {'A': 12, 'B': 70, 'C': 11, 'D': 10, 'E': 10, 'F': 10, 'G': 14, 'H': 12, 'I': 12})
hdr(dj, 1, 1, ['상품코드', '상품명', '판매가(원)', '결제수량', '환불수량', '판매수량', '판매합계(원)', '브랜드', '구분'], height=26)
for i, r_ in jm_raw.iterrows():
    x = 2 + i
    put(dj, f'A{x}', str(r_['상품코드'])); put(dj, f'B{x}', r_['상품명'])
    put(dj, f'C{x}', float(r_['판매가']), f_in, WON); put(dj, f'D{x}', float(r_['결제수량']), f_in, WON)
    put(dj, f'E{x}', float(r_['환불수량']), f_in, WON); put(dj, f'F{x}', float(r_['판매수량']), f_in, WON)
    put(dj, f'G{x}', float(r_['판매합계']), f_in, WON)
    put(dj, f'H{x}', f'=IF(AND(ISNUMBER(SEARCH("아넬라",B{x})),ISNUMBER(SEARCH("로지오가닉",B{x}))),"공통",IF(ISNUMBER(SEARCH("아넬라",B{x})),"아넬라",IF(ISNUMBER(SEARCH("로지오가닉",B{x})),"로지오가닉","기타")))')
    put(dj, f'I{x}', f'=IF(ISNUMBER(SEARCH("시크릿 특가",B{x})),"시크릿특가",IF(ISNUMBER(SEARCH("첫구매 전용",B{x})),"첫구매전용","그 외"))')
dj.freeze_panes = 'A2'
dj['A1'].comment = Comment('Source: 자사몰 10월 매출 CSV(상품 순위별 판매), 사용자 업로드, 2026-10-29. 기간 10/1~10/28, 판매합계는 실결제 기준(사용자 확인). 브랜드·구분은 상품명 기준 수식(아넬라/로지오가닉 포함 여부, "시크릿 특가"는 메타 시크릿링크 광고 유입 상품이라는 사용자 설명).', 'Claude')
dj['A1'].comment.width = 380; dj['A1'].comment.height = 130
def djr(col):
    return f"{S_DJ}!${col}$2:${col}$300"

jm = sheet('자사몰_매출', '70AD47')
widths(jm, {'A': 22, 'B': 13, 'C': 13, 'D': 13, 'E': 11, 'F': 13, 'G': 13, 'H': 12, 'I': 13, 'J': 12, 'K': 12})
put(jm, 'A1', '자사몰 10월 매출 (브랜드별) 과 메타 시크릿링크 광고', f_title)
put(jm, 'A2', '금액은 만원. 판매합계는 파일 값 그대로이며, 상품명에 "시크릿 특가"가 붙은 상품은 메타 시크릿링크 광고 유입 상품으로 간주. 기간은 10/1~10/28, 판매합계는 실결제 기준.', f_note)
jm.merge_cells('A2:K2'); jm.row_dimensions[2].height = 28
put(jm, 'A5', '브랜드별 자사몰 매출', f_sec)
hdr(jm, 6, 1, ['구분', '결제수량', '환불수량', '환불률', '판매수량', '판매합계', '시크릿특가 매출', '시크릿특가 비중', '첫구매 전용 매출', '첫구매 전용 비중', '그 외 매출'], height=32)
for i, b in enumerate(['아넬라', '로지오가닉', '공통']):
    r = 7 + i
    put(jm, f'A{r}', b if b != '공통' else '공통 (전품목 맛보기팩)', border=True)
    key = f'{djr("H")},"{b}"'
    for c, col in [('B', 'D'), ('C', 'E'), ('E', 'F')]:
        put(jm, f'{c}{r}', f'=SUMIFS({djr(col)},{key})', fmt=INT, border=True)
    put(jm, f'F{r}', f'=SUMIFS({djr("G")},{key})/10000', fmt=MAN, border=True)
    put(jm, f'G{r}', f'=SUMIFS({djr("G")},{key},{djr("I")},"시크릿특가")/10000', fmt=MAN, border=True)
    put(jm, f'I{r}', f'=SUMIFS({djr("G")},{key},{djr("I")},"첫구매전용")/10000', fmt=MAN, border=True)
put(jm, 'A10', '합계', f_bold, fill=fill_sub, border=True)
for c in 'BCEFGI':
    put(jm, f'{c}10', f'=SUM({c}7:{c}9)', f_bold, INT if c in 'BCE' else MAN, fill_sub, border=True)
for r in range(7, 11):
    bold = f_bold if r == 10 else f_norm
    fill = fill_sub if r == 10 else None
    put(jm, f'D{r}', f'=IF(B{r}=0,0,C{r}/B{r})', bold, '0.0%', fill, border=True)
    put(jm, f'H{r}', f'=IF(F{r}=0,0,G{r}/F{r})', bold, PCT, fill, border=True)
    put(jm, f'J{r}', f'=IF(F{r}=0,0,I{r}/F{r})', bold, PCT, fill, border=True)
    put(jm, f'K{r}', f'=F{r}-G{r}-I{r}', bold, MAN, fill, border=True)
put(jm, 'A11', '브랜드 구분은 상품명에 "아넬라"/"로지오가닉"이 있는지로 구분. 두 브랜드가 함께 적힌 상품(신규회원 한정 맛보기팩)은 공통으로 분리.', f_note)
jm.merge_cells('A11:K11')

put(jm, 'A13', '메타 시크릿링크 광고 vs 자사몰 시크릿특가 실매출', f_sec)
hdr(jm, 14, 1, ['브랜드', '시크릿링크 광고비', '메타 보고 매출', '시크릿특가 실매출', '보고 ÷ 실매출', '보고 ROAS', '실매출 ROAS', '메타 자사몰 전체 광고비', '자사몰 브랜드 전체 매출', '자사몰 매출 ÷ 전체 광고비'], height=44)
for i, b in enumerate(BRANDS):
    r = 15 + i; jr = 7 + i
    crit = f'{adr("A")},$A{r},{adr("B")},"메타 자사몰",{adr("C")},"*시크릿링크*"'
    put(jm, f'A{r}', b, f_bold, border=True)
    put(jm, f'B{r}', f'=SUMIFS({adr("E")},{crit})/10000', fmt=MAN, border=True)
    put(jm, f'C{r}', f'=SUMIFS({adr("I")},{crit})/10000', fmt=MAN, border=True)
    put(jm, f'D{r}', f'=G{jr}', fmt=MAN, border=True)
    put(jm, f'E{r}', f'=IF(D{r}=0,0,C{r}/D{r})', fmt='0.00"배"', border=True)
    put(jm, f'F{r}', f'=IF(B{r}=0,0,C{r}/B{r})', f_bold, PCT, border=True)
    put(jm, f'G{r}', f'=IF(B{r}=0,0,D{r}/B{r})', f_bold, PCT, border=True)
    put(jm, f'H{r}', f'=SUMIFS({adr("E")},{adr("A")},$A{r},{adr("B")},"메타 자사몰")/10000', fmt=MAN, border=True)
    put(jm, f'I{r}', f'=F{jr}', fmt=MAN, border=True)
    put(jm, f'J{r}', f'=IF(H{r}=0,0,I{r}/H{r})', fmt=PCT, border=True)
put(jm, 'A18', '시크릿링크 광고비는 캠페인명에 "시크릿링크"가 들어간 메타 자사몰 캠페인만 합산(아넬라 웰컴박스 제외). 메타 보고 매출이 실매출보다 큰 것은 귀속 기준 차이(조회 후 전환 등) 또는 환불·할인 반영 차이로 추정(판매합계는 실결제 기준).', f_note)
jm.merge_cells('A18:K18'); jm.row_dimensions[18].height = 40
put(jm, 'A20', '상품별 상위 (판매합계 순)', f_sec)
hdr(jm, 21, 1, ['순위', '상품명', '', '', '', '판매합계(만원)', '브랜드', '구분'], height=26)
jm.merge_cells('B21:E21')
top = jm_raw.sort_values('판매합계', ascending=False).head(10).reset_index()
for i in range(len(top)):
    r = 22 + i
    src = 2 + int(top.loc[i, 'index'])
    put(jm, f'A{r}', i + 1, al=CEN, border=True)
    put(jm, f'B{r}', f"={S_DJ}!B{src}", al=LEFT, border=True); jm.merge_cells(f'B{r}:E{r}')
    put(jm, f'F{r}', f"={S_DJ}!G{src}/10000", fmt=MAN, border=True)
    put(jm, f'G{r}', f"={S_DJ}!H{src}", al=CEN, border=True)
    put(jm, f'H{r}', f"={S_DJ}!I{src}", al=CEN, border=True)
jm.freeze_panes = 'A4'

# ---- wrap long note rows
def wrapnote(ws, ref, last, h=28):
    c = ws[ref]
    row = c.row
    ws.merge_cells(f'{c.column_letter}{row}:{last}{row}')
    c.alignment = LEFTTOP
    ws.row_dimensions[row].height = h
for ref in ['A2', 'A14', 'A26'] + [f'A{SM_R-2}', f'A{SM_R-1}', f'A{SM_R}']:
    wrapnote(sm, ref, 'K', 30)
for ref in ['A2', 'A27', 'A34', 'A41']:
    wrapnote(rs, ref, 'P', 28)
for ref in ['A2']:
    wrapnote(tg, ref, 'Q', 28)
    wrapnote(wk, ref, 'I', 28)
    wrapnote(ck, ref, 'G', 20)
    wrapnote(an, ref, 'H', 28)
    wrapnote(ep, ref, 'H', 20)
    wrapnote(ga, ref, 'F', 20)

# sheet order
order = ['요약', '10월_실적', '예산', '예산대비_집행', '자사몰_매출', '주차별_추이', '성과분석', '시간제한_검증', '목표_11-12월', '실행계획', '가정', '점검', '데이터_광고', '데이터_SS일별', '데이터_SS상품', '데이터_자사몰']
wb._sheets = [wb[n] for n in order]
for ws in wb.worksheets:
    ws.sheet_view.zoomScale = 100
wb.save(OUT + 'report.xlsx')
print('saved', AN_LAST, SM_R, TGR['11월']['judge'], TGR['12월']['judge'], ssr)
